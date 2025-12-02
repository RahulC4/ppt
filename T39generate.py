# ============================================================
# generate_ppt.py – FINAL STABLE CORPORATE TEMPLATE VERSION
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import PP_PLACEHOLDER as PH

from utils import get_env, logger, now_ts, ensure_dir, text_client, image_client
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob


# ------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------
ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
CORPORATE_TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "corporate.pptx")


# ------------------------------------------------------------
# BASIC HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    return int(match.group(1)) if match else None


def fallback_plan(n):
    slides = []
    for i in range(n):
        slides.append({
            "title": f"Slide {i+1}",
            "bullets": [
                f"Key takeaway {i+1}",
                f"Supporting point {i+1}.1",
                f"Supporting point {i+1}.2",
            ],
        })
    return slides


# ------------------------------------------------------------
# LLM
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    references_text = references_text or []

    sys_prompt = (
        "Return EXACTLY this format:\n\n"
        "Slide 1: Title\n- Bullet\n- Bullet\n- Bullet\n\n"
        f"You MUST create exactly {num_slides} slides.\n"
        "NO JSON. NO commentary.\n\n"
        f"References:\n{' '.join(references_text)[:2500]}"
    )

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": prompt},
            ],
            max_completion_tokens=1400,
            temperature=0.7,
        )
        return parse_text_plan(resp.choices[0].message.content, num_slides)

    except Exception as e:
        logger.warning(f"LLM failed → fallback used: {e}")
        return fallback_plan(num_slides)


def parse_text_plan(text, num_slides):
    blocks = re.split(r"\n(?=Slide\s+\d+:)", text)
    slides = []

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        title = lines[0].split(":", 1)[-1].strip()
        bullets = [l[1:].strip() for l in lines[1:] if l.startswith("-")]

        while len(bullets) < 3:
            bullets.append("Additional point")

        slides.append({"title": title, "bullets": bullets[:7]})

    return slides if len(slides) >= num_slides else None


# ------------------------------------------------------------
# IMAGE
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " minimal corporate illustration no text",
            size="1024x1024",
        )
        img_bytes = base64.b64decode(resp.data[0].b64_json)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name
    except:
        logger.exception("Image failed")
        return None


# ------------------------------------------------------------
# SAFE PLACEHOLDER WRITERS
# ------------------------------------------------------------
def set_title(slide, text):
    for ph in slide.placeholders:
        if ph.placeholder_format.type in (PH.TITLE, PH.CENTER_TITLE):
            ph.text = text
            return


def clear_and_set_bullets(slide, bullets):
    body = None
    for ph in slide.placeholders:
        if ph.placeholder_format.type == PH.BODY and ph.has_text_frame:
            body = ph
            break

    if body is None:
        return

    tf = body.text_frame
    tf.clear()

    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(20)


def add_image(slide, img_path, body_shape=None):
    try:
        if body_shape:
            body_shape.width = int(body_shape.width * 0.6)
            left = body_shape.left + body_shape.width + Inches(0.2)
            top = body_shape.top
        else:
            left = Inches(7)
            top = Inches(1.5)

        slide.shapes.add_picture(img_path, left, top, width=Inches(2.8))
    except:
        logger.exception("Image placement failed")


def remove_speaker_text(slide):
    for ph in slide.placeholders:
        if ph.has_text_frame and "speaker" in ph.text.lower():
            ph.text = ""


def update_date(slide):
    current = datetime.now().strftime("%B %Y")
    for ph in slide.placeholders:
        if ph.has_text_frame and re.search(r"\b20\d{2}\b", ph.text):
            ph.text = current


def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slide_id = xml_slides[index].rId
    prs.part.drop_rel(slide_id)
    del xml_slides[index]


# ------------------------------------------------------------
# CORPORATE BUILDER — FINAL FIXED
# ------------------------------------------------------------
def build_corporate_ppt(plan, image_required):

    prs = Presentation(CORPORATE_TEMPLATE_PATH)

    title_layout = prs.slide_layouts[0]
    agenda_layout = prs.slide_layouts[1]
    content_layout = prs.slide_layouts[1]
    thanks_layout = prs.slide_layouts[-1]

    for i in reversed(range(len(prs.slides))):
        delete_slide(prs, i)

    # ---------- SLIDE 1 ----------
    title_slide = prs.slides.add_slide(title_layout)
    set_title(title_slide, plan[0]["title"])
    remove_speaker_text(title_slide)
    update_date(title_slide)

    # ---------- SLIDE 2 (AGENDA - REPLACED) ----------
    agenda_slide = prs.slides.add_slide(agenda_layout)
    set_title(agenda_slide, "Agenda")
    agenda_items = [s["title"] for s in plan]
    clear_and_set_bullets(agenda_slide, agenda_items)

    # ---------- SLIDES 3 → N (ALWAYS WITH TITLES + BULLETS) ----------
    for sp in plan:
        slide = prs.slides.add_slide(content_layout)
        set_title(slide, sp["title"])
        clear_and_set_bullets(slide, sp["bullets"])

        if image_required:
            img_path = generate_visual_image(sp["title"])
            if img_path:
                body = next(ph for ph in slide.placeholders if ph.placeholder_format.type == PH.BODY)
                add_image(slide, img_path, body)

    # ---------- FINAL THANK YOU ----------
    prs.slides.add_slide(thanks_layout)

    out_path = os.path.join(tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# FINAL PIPELINE
# ------------------------------------------------------------
def generate_presentation(prompt, requested_num_slides=5, tag_filters=None, template_style=None, image_required=False):

    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []
    if not refs:
        return None, {"error": True, "message": "No relevant content found."}

    reference_text = [(r.get("text") or "")[:500] for r in refs]

    num_slides = requested_num_slides or parse_user_intent(prompt) or 5

    plan = call_llm_plan(prompt, reference_text, num_slides)

    if not plan:
        return None, {"error": True, "message": "Insufficient content."}

    ppt_path = build_corporate_ppt(plan, image_required)

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(ppt_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(plan) + 3,
        "ppt_file": fname,
        "image_required": image_required,
        "template_style": template_style,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode(), f"logs/{fname}.json")

    return ppt_path, log
