# ============================================================
# generate_ppt.py – NO JSON | Images Checkbox | Safe Scaling | AUTO SIZE FIX | SAFE PLACEHOLDERS
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE

from utils import (
    get_env, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
CORPORATE_TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "corporate.pptx")

# ------------------------------------------------------------
# SAFE BODY PLACEHOLDER FETCH (CRASH FIX)
# ------------------------------------------------------------
def get_safe_body_placeholder(slide):
    try:
        return slide.placeholders[1]
    except Exception:
        for shp in slide.shapes:
            if getattr(shp, "has_text_frame", False) and shp is not slide.shapes.title:
                return shp
    return None


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    return int(match.group(1)) if match else None


def fallback_plan(n):
    return [
        {
            "title": f"Slide {i+1}",
            "bullets": [
                f"Key takeaway for slide {i+1}",
                f"Supporting detail {i+1}.1",
                f"Supporting detail {i+1}.2",
            ],
        } for i in range(n)
    ]


# ------------------------------------------------------------
# LLM TEXT PLAN (NO JSON)
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    references_text = references_text or []

    sys_prompt = (
        "You are a professional presentation creator.\n\n"
        "Return ONLY in this format:\n\n"
        "Slide 1: Title\n"
        "- Bullet\n"
        "- Bullet\n"
        "- Bullet\n\n"
        f"You MUST create exactly {num_slides} slides.\n"
        "Do NOT return JSON.\n"
        "Do NOT explain anything.\n\n"
        f"Reference:\n{' '.join(references_text)[:2500]}"
    )

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": prompt},
            ],
            max_completion_tokens=1400,
            temperature=1,
        )

        raw = resp.choices[0].message.content.strip()
        return parse_text_plan(raw, num_slides)

    except Exception as e:
        logger.warning(f"LLM failed → fallback used: {e}")
        return fallback_plan(num_slides)


def parse_text_plan(text, num_slides):
    slides = []
    blocks = re.split(r"\n(?=Slide \d+:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        title = lines[0].split(":", 1)[-1].strip()
        bullets = [l.replace("-", "").strip() for l in lines[1:] if l.startswith("-")]

        if len(bullets) < 3:
            bullets += [f"Additional point {i+1}" for i in range(3 - len(bullets))]

        slides.append({"title": title, "bullets": bullets[:6]})

    return slides[:num_slides] if len(slides) >= num_slides else None


# ------------------------------------------------------------
# IMAGE GENERATION
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " Professional minimal illustration. No text.",
            size="1024x1024",
        )
        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(base64.b64decode(b64))
        tmp.close()
        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# ✅✅✅ SAFE PPT BUILDER (NO TEXT SHRINK + NO PLACEHOLDER CRASH)
# ------------------------------------------------------------
def build_default_ppt(slides, image_required):
    prs = Presentation()

    for sp in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = sp["title"]

        body = get_safe_body_placeholder(slide)

        if body is None:
            body = slide.shapes.add_textbox(
                Inches(0.5), Inches(2), prs.slide_width - Inches(1), Inches(4)
            )

        tf = body.text_frame
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.word_wrap = True
        tf.clear()

        for b in sp["bullets"]:
            p = tf.add_paragraph()
            p.text = b
            p.level = 0
            p.font.size = Pt(20)

        body.top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.3)

        if image_required and sp.get("image_path"):
            body.width = prs.slide_width - Inches(4)
            body.left = Inches(0.5)

            slide.shapes.add_picture(
                sp["image_path"],
                prs.slide_width - Inches(3.5),
                body.top,
                width=Inches(3),
            )
        else:
            body.left = Inches(0.5)
            body.width = prs.slide_width - Inches(1)

    out_path = os.path.join(tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# ✅ FINAL PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt,
    requested_num_slides=5,
    tag_filters=None,
    template_style=None,
    image_required=False,
):
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []
    if not refs:
        return None, {"error": True, "message": "No matching content found in sample PPTs."}

    reference_text = [(r.get("text") or "")[:500] for r in refs]

    detected_slides = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected_slides or 5

    plan = call_llm_plan(prompt, reference_text, num_slides)

    if not plan:
        return None, {
            "error": True,
            "message": "Not enough relevant content. Try fewer slides or rephrase."
        }

    slides = []
    for sp in plan:
        img_path = generate_visual_image(sp["title"]) if image_required else None
        slides.append({
            "title": sp["title"],
            "bullets": sp["bullets"],
            "image_path": img_path,
        })

    out_path = build_default_ppt(slides, image_required)

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides),
        "ppt_file": fname,
        "image_required": image_required,
        "template_style": template_style,
        "error": False,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode(), f"logs/{fname}.json")
    return out_path, log
