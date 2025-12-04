# ============================================================
# generate_ppt.py – NO JSON | Images Checkbox | Corporate Template (FIXED + IMAGE RULES)
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from utils import (
    get_env,
    logger,
    now_ts,
    ensure_dir,
    text_client,
    image_client,
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
CORPORATE_TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "corporate.pptx")


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    return int(match.group(1)) if match else None


def fallback_plan(n):
    return [
        {
            "title": f"Slide {i+1}",
            "bullets": [
                f"Main point {i+1}",
                f"Supporting detail {i+1}.1",
                f"Supporting detail {i+1}.2",
            ],
        }
        for i in range(n)
    ]


# ------------------------------------------------------------
# LLM PLAN
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    references_text = references_text or []

    sys_prompt = (
        "Return ONLY this format:\n"
        "Slide 1: Title\n"
        "- Bullet\n"
        "- Bullet\n"
        "- Bullet\n\n"
        f"Create EXACTLY {num_slides} slides."
    )

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": prompt},
            ],
            max_completion_tokens=1400,
            temperature=1,
        )

        raw = resp.choices[0].message.content.strip()
        return parse_text_plan(raw, num_slides)

    except Exception as e:
        logger.warning(f"LLM failed → fallback used: {e}")
        return fallback_plan(num_slides)


def parse_text_plan(text, num_slides):
    slides = []
    blocks = re.split(r"\n(?=Slide\s+\d+\s*:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        title = lines[0].split(":", 1)[-1].strip()
        bullets = [l[1:].strip() for l in lines[1:] if l.startswith("-")]
        bullets = bullets[:8]

        slides.append({"title": title, "bullets": bullets})

    return slides[:num_slides]


# ------------------------------------------------------------
# GENERIC HELPERS
# ------------------------------------------------------------
def set_title(slide, text):
    for shp in slide.shapes:
        if shp.has_text_frame:
            shp.text = text
            break


def update_date(slide):
    current = datetime.now().strftime("%B %Y")
    for shp in slide.shapes:
        if shp.has_text_frame:
            if re.search(r"\b20\d{2}\b", shp.text):
                shp.text = current


def add_bullet_textbox(slide, bullets, left, top, width, height, font_size=20,color=RGBColor(0,0,0)):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.clear()

    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = color

    return tb


def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slide_id = xml_slides[index].rId
    prs.part.drop_rel(slide_id)
    del xml_slides[index]


# ------------------------------------------------------------
# IMAGE HELPERS
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " Professional minimal illustration. No text.",
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


def add_image_to_right(slide, body_shape, img_path: str):
    try:
        left = body_shape.left + body_shape.width + Inches(0.15)
        top = body_shape.top
        slide.shapes.add_picture(img_path, left, top, width=Inches(2.8))
    except Exception:
        logger.exception("Image placement failed")


# ------------------------------------------------------------
# CORPORATE BUILDER WITH FINAL IMAGE RULES
# ------------------------------------------------------------
def build_corporate_ppt(plan, image_required=False):
    prs = Presentation(CORPORATE_TEMPLATE_PATH)

    title_layout = prs.slides[0].slide_layout
    agenda_layout = prs.slides[1].slide_layout
    content_layout = prs.slides[2].slide_layout
    thankyou_layout = prs.slides[-1].slide_layout

    for i in reversed(range(len(prs.slides))):
        delete_slide(prs, i)

    # ---- 1) TITLE SLIDE (NO IMAGE) ----
    title_slide = prs.slides.add_slide(title_layout)
    set_title(title_slide, plan[0]["title"])
    update_date(title_slide)

    # ---- 2) AGENDA SLIDE (IMAGE ALLOWED ✅) ----
    agenda_slide = prs.slides.add_slide(agenda_layout)
    fill = agenda_slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(242, 244, 247)
    set_title(agenda_slide, "Agenda")

    agenda_items = [s["title"] for s in plan]

    body = add_bullet_textbox(
        agenda_slide,
        agenda_items,
        left=Inches(1),
        top=Inches(2),
        width=prs.slide_width - Inches(4.5) if image_required else prs.slide_width - Inches(2),
        height=Inches(4),
        font_size=32,
        color = RGBColor(0,102,204)
    )

    if image_required:
        img = generate_visual_image("Agenda presentation illustration")
        if img:
            add_image_to_right(agenda_slide, body, img)

    # ---- CONTENT SLIDES (IMAGE ALLOWED ✅) ----
    for sp in plan:
        slide = prs.slides.add_slide(content_layout)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(242, 244, 247)
        set_title(slide, sp["title"])

        body = add_bullet_textbox(
            slide,
            sp["bullets"],
            left=Inches(1),
            top=Inches(2),
            width=prs.slide_width - Inches(4.5) if image_required else prs.slide_width - Inches(2),
            height=Inches(4.5),
            font_size=20,
        )

        if image_required:
            img = generate_visual_image(sp["title"])
            if img:
                add_image_to_right(slide, body, img)

    # ---- THANK YOU SLIDE (NO IMAGE) ----
    prs.slides.add_slide(thankyou_layout)

    out = os.path.join(tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out)
    return out


# ------------------------------------------------------------
# MAIN PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt,
    requested_num_slides=5,
    tag_filters=None,
    template_style=None,
    image_required=False,
):
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []
    reference_text = [(r.get("text") or "")[:500] for r in refs]

    detected = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected or 5

    plan = call_llm_plan(prompt, reference_text, num_slides)

    ppt_path = build_corporate_ppt(plan, image_required=image_required)
    total_slides = len(plan) + 3

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(ppt_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": total_slides,
        "ppt_file": fname,
        "image_required": image_required,
        "template_style": template_style,
        "error": False,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode("utf-8"), f"logs/{fname}.json")
    return ppt_path, log
