# ============================================================
# generate_ppt.py – NO JSON | Images Checkbox | Corporate Template
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import PP_PLACEHOLDER as PH
from PIL import Image

from utils import (
    get_env,
    logger,
    now_ts,
    ensure_dir,
    text_client,
    image_client,
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

# ------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------
ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
CORPORATE_TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "corporate.pptx")


# ------------------------------------------------------------
# BASIC HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    """Try to detect 'N slides' from user prompt."""
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        return int(match.group(1))
    return None


def fallback_plan(n):
    """Deterministic fallback slide plan."""
    slides = []
    for i in range(n):
        slides.append(
            {
                "title": f"Slide {i + 1}",
                "bullets": [
                    f"Key takeaway for slide {i + 1}",
                    f"Supporting detail {i + 1}.1",
                    f"Supporting detail {i + 1}.2",
                ],
            }
        )
    return slides


# ------------------------------------------------------------
# ✅ TEXT-BASED LLM (NO JSON)
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    """
    Ask GPT to create a slide plan in plain text, NOT JSON.

    Expected format in the response:

        Slide 1: Title
        - Bullet
        - Bullet

        Slide 2: Another title
        - Bullet
        ...

    Then we parse this with parse_text_plan().
    """
    references_text = references_text or []

    sys_prompt = (
        "You are a professional presentation creator.\n\n"
        "Return the slide plan in this EXACT TEXT format only:\n\n"
        "Slide 1: Title of the slide\n"
        "- Bullet point one\n"
        "- Bullet point two\n"
        "- Bullet point three\n\n"
        "Slide 2: Title of the second slide\n"
        "- Bullet point one\n"
        "- Bullet point two\n\n"
        f"You MUST create exactly {num_slides} slides.\n"
        "Do NOT return JSON.\n"
        "Do NOT add explanations or commentary.\n\n"
        "Use the reference content for facts and wording when possible:\n"
        f"{' '.join(references_text)[:2500]}"
    )

    user_prompt = f"Create a professional presentation for: {prompt}"

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt},
            ],
            max_completion_tokens=1400,
            temperature=0.7,
        )

        raw_text = resp.choices[0].message.content.strip()
        return parse_text_plan(raw_text, num_slides)

    except Exception as e:
        logger.warning(f"LLM failed → fallback used: {e}")
        return fallback_plan(num_slides)


# ------------------------------------------------------------
# ✅ TEXT → STRUCTURED SLIDES
# ------------------------------------------------------------
def parse_text_plan(text, num_slides):
    """
    Parse the plain-text plan into a list of:
        { "title": str, "bullets": [str] }
    """
    slides = []

    # Split on lines that start a new slide: "Slide 1:", "Slide 2:", etc.
    blocks = re.split(r"\n(?=Slide\s+\d+\s*:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        # First line: "Slide X: Title"
        title_line = lines[0]
        if ":" in title_line:
            title = title_line.split(":", 1)[1].strip()
        else:
            title = title_line

        bullets = []
        for l in lines[1:]:
            if l.startswith("-"):
                bullets.append(l[1:].strip())

        # Enforce minimum bullets per slide
        if len(bullets) < 3:
            bullets += [f"Additional point {i + 1}" for i in range(3 - len(bullets))]

        slides.append(
            {
                "title": title,
                "bullets": bullets[:8],  # cap to avoid overflow
            }
        )

    # Enforce slide count strictly – if not enough, we signal failure
    if len(slides) < num_slides:
        logger.warning(
            f"Parsed only {len(slides)} slides but {num_slides} were requested."
        )
        return None

    return slides[:num_slides]


# ------------------------------------------------------------
# IMAGE GENERATION (OPTIONAL)
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    """Generate an image file from GPT image model. Returns a temp file path or None."""
    if not prompt:
        return None

    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " Professional minimal illustration. No text.",
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# SLIDE UTILS (TITLE / BULLETS / IMAGES / DATES)
# ------------------------------------------------------------
def _set_title_text(slide, text: str):
    """Set the slide's title text in a safe way."""
    try:
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            slide.shapes.title.text = text
            return
    except Exception:
        pass

    # Fallback: any TITLE / CENTER_TITLE / SUBTITLE placeholder
    for ph in slide.placeholders:
        if ph.placeholder_format.type in (PH.TITLE, PH.CENTER_TITLE, PH.SUBTITLE):
            if getattr(ph, "has_text_frame", False):
                ph.text = text
                return

    # As a last resort, do nothing (avoid crashing).


def _set_bullets(slide, bullets):
    """Fill BODY text placeholder with bullet points."""
    body_shape = None

    # Prefer BODY placeholder
    for ph in slide.placeholders:
        if ph.placeholder_format.type == PH.BODY and getattr(
            ph, "has_text_frame", False
        ):
            body_shape = ph
            break

    # Fallback: any text frame that isn't the title
    if body_shape is None:
        for shp in slide.shapes:
            if getattr(shp, "has_text_frame", False) and shp is not slide.shapes.title:
                body_shape = shp
                break

    if body_shape is None:
        logger.warning("No suitable body placeholder found on a slide; skipping bullets")
        return

    tf = body_shape.text_frame
    tf.clear()

    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(20)


def _add_image_to_slide(slide, img_path: str):
    """Place image on the right side, shrinking body if we can."""
    try:
        body_shape = None
        for ph in slide.placeholders:
            if ph.placeholder_format.type == PH.BODY:
                body_shape = ph
                break

        if body_shape is not None:
            # shrink body width to leave space for image
            body_shape.width = int(body_shape.width * 0.55)
            left_img = body_shape.left + body_shape.width + Inches(0.3)
            top_img = body_shape.top
        else:
            # generic position on the right
            left_img = Inches(7)
            top_img = Inches(1.5)

        slide.shapes.add_picture(img_path, left_img, top_img, height=Inches(3))
    except Exception:
        logger.exception("Image placement failed")


def _update_date_placeholders(slide):
    """Replace any text containing a year with current Month YYYY."""
    current = datetime.now().strftime("%B %Y")
    for shape in slide.placeholders:
        if getattr(shape, "has_text_frame", False):
            txt = shape.text or ""
            if re.search(r"\b20\d{2}\b", txt):
                shape.text = current


def _delete_slide(prs, index: int):
    """Delete slide at index using python-pptx internals."""
    xml_slides = prs.slides._sldIdLst  # type: ignore[attr-defined]
    slide_id = xml_slides[index].rId
    prs.part.drop_rel(slide_id)
    del xml_slides[index]


# ------------------------------------------------------------
# DEFAULT PPT BUILDER (NO TEMPLATE)
# ------------------------------------------------------------
def build_default_ppt(plan, image_required: bool):
    prs = Presentation()

    for sp in plan:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        _set_title_text(slide, sp["title"])
        _set_bullets(slide, sp["bullets"])

        if image_required:
            img_path = generate_visual_image(sp["title"])
            if img_path:
                _add_image_to_slide(slide, img_path)

    out_path = os.path.join(
        tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx"
    )
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# CORPORATE TEMPLATE BUILDER
# ------------------------------------------------------------
def build_corporate_ppt(plan, image_required: bool):
    """
    Build deck using templates/corporate.pptx.

    Structure:
      1. Title slide (from template layout, text + current date)
      2. Agenda slide (bullets = all section titles)
      3..N+2. Content slides (from plan, optional images)
      Last. Thank-you slide (unchanged layout; no image)
    """
    if not os.path.exists(CORPORATE_TEMPLATE_PATH):
        logger.warning(
            "Corporate template not found at %s, falling back to default layout",
            CORPORATE_TEMPLATE_PATH,
        )
        return build_default_ppt(plan, image_required)

    prs = Presentation(CORPORATE_TEMPLATE_PATH)

    if len(prs.slides) == 0:
        logger.warning("Corporate template has no slides; falling back to default.")
        return build_default_ppt(plan, image_required)

    # Capture layouts used by first, second and last slide
    title_layout = prs.slides[0].slide_layout
    agenda_layout = prs.slides[1].slide_layout if len(prs.slides) > 1 else prs.slide_layouts[1]
    thanks_layout = prs.slides[-1].slide_layout

    # Remove ALL existing slides from this presentation
    for idx in reversed(range(len(prs.slides))):
        _delete_slide(prs, idx)

    # ---- Title slide ----
    deck_title = plan[0]["title"] if plan and plan[0].get("title") else "Executive Summary"
    title_slide = prs.slides.add_slide(title_layout)
    _set_title_text(title_slide, deck_title)
    _update_date_placeholders(title_slide)

    # ---- Agenda slide ----
    agenda_slide = prs.slides.add_slide(agenda_layout)
    _set_title_text(agenda_slide, "Agenda")

    agenda_items = [s["title"] for s in plan if s.get("title")]
    if not agenda_items:
        agenda_items = [f"Section {i + 1}" for i in range(len(plan))]
    _set_bullets(agenda_slide, agenda_items)

    # ---- Content slides (no template sample slides kept) ----
    for sp in plan:
        slide = prs.slides.add_slide(agenda_layout)
        _set_title_text(slide, sp["title"])
        _set_bullets(slide, sp["bullets"])

        # images ONLY on these content slides
        if image_required:
            img_path = generate_visual_image(sp["title"])
            if img_path:
                _add_image_to_slide(slide, img_path)

    # ---- Thank-you slide at the end (no image) ----
    prs.slides.add_slide(thanks_layout)

    out_path = os.path.join(
        tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx"
    )
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# ✅ FINAL PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt,
    requested_num_slides=5,
    tag_filters=None,
    template_style=None,
    image_required=False,
):
    """
    Main entry used by app.py.

    - Semantic search in Chroma
    - Ask LLM for a slide plan (no JSON)
    - Build PPT either with default layout or corporate template
    - Return (ppt_path, log_dict)
    """
    # 1) Semantic search over uploaded PPTs
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []
    if not refs:
        return None, {
            "error": True,
            "message": "I couldn’t find relevant content in your sample PPTs. "
            "Try a prompt closer to your uploaded decks.",
        }

    reference_text = [(r.get("text") or "")[:500] for r in refs]

    # 2) Decide slide count (user override > prompt hint > default 5)
    detected_slides = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected_slides or 5

    # 3) LLM plan (NO JSON)
    plan = call_llm_plan(prompt, references_text=reference_text, num_slides=num_slides)

    if not plan:
        return None, {
            "error": True,
            "message": "Not enough relevant content to generate this many slides. "
            "Try using fewer slides or rephrasing your prompt.",
        }

    # 4) Build PPT (corporate template vs default)
    template_key = (template_style or "").lower()

    if template_key == "corporate":
        ppt_path = build_corporate_ppt(plan, image_required=image_required)
        total_slides = len(plan) + 3  # title + agenda + thank-you
    else:
        ppt_path = build_default_ppt(plan, image_required=image_required)
        total_slides = len(plan)

    # 5) Upload + log
    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(ppt_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": total_slides,
        "ppt_file": fname,
        "error": False,
        "image_required": image_required,
        "template_style": template_style,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode("utf-8"), f"logs/{fname}.json")

    return ppt_path, log
