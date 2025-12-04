title_tf = agenda_slide.shapes.title.text_frame
p = title_tf.paragraphs[0]
p.font.size = Pt(32)
p.font.color.rgb = RGBColor(0, 102, 204)  # Corporate Blue
p.alignment = PP_ALIGN.LEFT

title_tf = slide.shapes.title.text_frame
p = title_tf.paragraphs[0]
p.font.size = Pt(32)
p.font.color.rgb = RGBColor(0, 102, 204)  # Corporate Blue
p.alignment = PP_ALIGN.LEFT

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

p.font.bold = True
