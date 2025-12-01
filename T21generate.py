# ============================================================
# generate_ppt.py – Semantic + Fixed Layout + UI Friendly
# ✅ WITH GUARANTEED EVEN-SPLIT FALLBACK
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64

from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image

from utils import (
    get_env, safe_json_load, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

# ------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------
ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        return int(match.group(1))
    return None


def chunk_list(lst, n):
    """Evenly split list into n chunks"""
    k, m = divmod(len(lst), n)
    return [
        lst[i * k + min(i, m):(i + 1) * k + min(i + 1, m)]
        for i in range(n)
    ]


# ------------------------------------------------------------
# ✅ SAFE LLM PLAN GENERATOR
# ------------------------------------------------------------
def call_llm_plan(
    prompt,
    style,
    design_context=None,
    references_text=None,
    num_slides=None,
    theme=None,
    text_density=None,
):
    references_text = references_text or []

    sys_prompt = (
        "You are a presentation planner.\n"
        "Return STRICT JSON ONLY in this exact format:\n"
        "[{\"title\": str, \"bullets\": [str], "
        "\"visual_required\": bool, \"visual_prompt\": str }]\n"
        "If the user asks for images, set visual_required=true.\n"
        "Do NOT put any words or labels inside visual_prompt images.\n\n"
        "You may use these reference snippets as content guidance:\n"
        f"{json.dumps(references_text)[:2500]}"
    )

    user_prompt = f"Create a professional presentation plan for: {prompt}"
    if num_slides:
        user_prompt += f" Use exactly {num_slides} slides."

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt},
            ],
            max_completion_tokens=1200,
            temperature=1,
        )

        raw = resp.choices[0].message.content
        plan = safe_json_load(raw)

        if not isinstance(plan, list) or len(plan) == 0:
            raise ValueError("Invalid plan JSON")

        return plan

    except Exception as e:
        logger.warning(f"LLM failed, switching to EVEN SPLIT fallback: {e}")
        return None  # Let main pipeline do deterministic split


# ------------------------------------------------------------
# IMAGE GENERATION
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    if not prompt:
        return None

    img_prompt = prompt + " Minimal, professional illustration. No text labels."

    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=img_prompt,
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# PPT BUILDER
# ------------------------------------------------------------
def build_ppt(slides):
    prs = Presentation()

    for sp in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])

        slide.shapes.title.text = sp.get("title", "")

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        for b in sp.get("bullets", []):
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(20)

        body.top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.3)

        img_path = sp.get("image_path")

        if not img_path:
            body.left = Inches(0.5)
            body.width = prs.slide_width - Inches(1.0)
            continue

        body.left = Inches(0.5)
        body.width = prs.slide_width - Inches(4.0)

        try:
            img = Image.open(img_path)
            w, h = img.size
            aspect = w / h if h else 1.0

            max_w = Inches(3.0)
            max_h = Inches(2.5)

            if aspect >= 1:
                final_w = max_w
                final_h = final_w / aspect
            else:
                final_h = max_h
                final_w = final_h * aspect

            left = prs.slide_width - final_w - Inches(0.5)
            top = body.top

            slide.shapes.add_picture(img_path, left, top, width=final_w, height=final_h)
        except Exception:
            logger.exception("Image placement failed")

    out_path = os.path.join(
        tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx"
    )
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# ✅ MAIN PIPELINE WITH GUARANTEED SPLIT BACKUP
# ------------------------------------------------------------
def generate_presentation(
    prompt: str,
    style="Auto",
    requested_num_slides=None,
    theme=None,
    text_density=None,
    tag_filters=None,
    template_style=None,
):

    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []

    if not refs:
        msg = "No matching content found in sample PPTs."
        return None, {"error": True, "message": msg}

    reference_text = []
    for r in refs:
        snippet = (r.get("text") or "")[:500]
        if snippet:
            reference_text.append(snippet)

    detected_slides = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected_slides or 5

    # ✅ Try LLM first
    plan = call_llm_plan(
        prompt=prompt,
        style=style,
        references_text=reference_text,
        num_slides=num_slides,
        theme=theme,
        text_density=text_density,
    )

    # ✅ FALLBACK: EVEN SPLIT IF LLM FAILS
    if not isinstance(plan, list) or len(plan) != num_slides:
        logger.warning("Using EVEN SPLIT fallback for slides")

        chunks = chunk_list(reference_text, num_slides)

        plan = []
        for i, chunk in enumerate(chunks):
            bullets = []
            for c in chunk:
                bullets.extend([x.strip() for x in c.split(".") if x.strip()][:5])

            plan.append({
                "title": f"Slide {i+1}",
                "bullets": bullets[:5],
                "visual_required": False,
                "visual_prompt": "",
            })

    force_images = "image" in prompt.lower() or "images" in prompt.lower()
    slides = []

    for sp in plan:
        img_path = None
        if sp.get("visual_required") or force_images:
            img_path = generate_visual_image(sp.get("visual_prompt"))

        slides.append({
            "title": sp.get("title", "Untitled"),
            "bullets": sp.get("bullets", []),
            "image_path": img_path,
        })

    out_path = build_ppt(slides)

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides),
        "ppt_file": fname,
        "error": False,
        "template_style": template_style,
    }

    upload_json_to_blob(
        json.dumps(log, indent=2).encode("utf-8"),
        f"logs/{fname}.json",
    )

    return out_path, log
