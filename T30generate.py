# ============================================================
# generate_ppt.py – NO JSON | Text-Based Stable Slide Generator
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64

from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image

from utils import (
    get_env, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

# ------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------
ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        return int(match.group(1))
    return None


# ------------------------------------------------------------
# ✅ ✅ NEW TEXT-BASED LLM PLAN (NO JSON)
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    references_text = references_text or []

    sys_prompt = (
        "You are a professional presentation creator.\n\n"
        "Return the slide plan in the following STRICT TEXT format ONLY:\n\n"
        "Slide 1: Title\n"
        "- Bullet 1\n"
        "- Bullet 2\n\n"
        "Slide 2: Title\n"
        "- Bullet 1\n"
        "- Bullet 2\n\n"
        "Rules:\n"
        f"- You MUST create exactly {num_slides} slides.\n"
        "- DO NOT return JSON.\n"
        "- DO NOT add explanations.\n"
        "- ONLY return formatted slide text.\n\n"
        "Use this content as reference:\n"
        f"{' '.join(references_text)[:2500]}"
    )

    user_prompt = f"Create a professional presentation for: {prompt}"

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt},
            ],
            max_completion_tokens=1200,
            temperature=0.7,
        )

        raw_text = resp.choices[0].message.content.strip()
        return parse_text_plan(raw_text, num_slides)

    except Exception as e:
        logger.warning(f"LLM failed, using fallback: {e}")
        return fallback_plan(num_slides)


# ------------------------------------------------------------
# ✅ ✅ TEXT → STRUCTURED SLIDES
# ------------------------------------------------------------
def parse_text_plan(text, num_slides):
    slides = []

    blocks = re.split(r"\n(?=Slide \d+:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        title_line = lines[0]
        title = title_line.replace("Slide", "").split(":", 1)[-1].strip()

        bullets = []
        for l in lines[1:]:
            if l.startswith("-"):
                bullets.append(l.replace("-", "").strip())

        slides.append(
            {
                "title": title,
                "bullets": bullets or ["Key point"],
                "visual_required": False,
                "visual_prompt": "",
            }
        )

    # ✅ FORCE EXACT SLIDE COUNT
    if len(slides) > num_slides:
        slides = slides[:num_slides]
    while len(slides) < num_slides:
        slides.append(
            {
                "title": f"Slide {len(slides)+1}",
                "bullets": ["Key point"],
                "visual_required": False,
                "visual_prompt": "",
            }
        )

    return slides


def fallback_plan(n):
    slides = []
    for i in range(n):
        slides.append(
            {
                "title": f"Slide {i+1}",
                "bullets": [
                    f"Key point 1 for slide {i+1}",
                    f"Key point 2 for slide {i+1}",
                ],
                "visual_required": False,
                "visual_prompt": "",
            }
        )
    return slides


# ------------------------------------------------------------
# IMAGE GENERATION
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    if not prompt:
        return None

    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " Minimal professional illustration. No text.",
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# PPT BUILDER
# ------------------------------------------------------------
def build_ppt(slides):
    prs = Presentation()

    for sp in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = sp.get("title", "")

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        for b in sp.get("bullets", []):
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(20)

    out_path = os.path.join(
        tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx"
    )
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# ✅ ✅ MAIN PIPELINE (NO JSON DEPENDENCY)
# ------------------------------------------------------------
def generate_presentation(
    prompt: str,
    style="Auto",
    requested_num_slides=None,
    theme=None,
    tag_filters=None,
    template_style=None,
):
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []

    if not refs:
        msg = "No relevant content found in uploaded PPTs."
        return None, {"error": True, "message": msg}

    reference_text = []
    for r in refs:
        snippet = (r.get("text") or "")[:500]
        if snippet:
            reference_text.append(snippet)

    detected_slides = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected_slides or 5

    plan = call_llm_plan(
        prompt=prompt,
        references_text=reference_text,
        num_slides=num_slides,
    )

    slides = []
    for sp in plan:
        slides.append(
            {
                "title": sp["title"],
                "bullets": sp["bullets"],
                "image_path": None,
            }
        )

    out_path = build_ppt(slides)

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides),
        "ppt_file": fname,
        "error": False,
        "template_style": template_style,
    }

    upload_json_to_blob(
        json.dumps(log, indent=2).encode("utf-8"),
        f"logs/{fname}.json",
    )

    return out_path, log
