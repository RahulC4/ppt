# ============================================================
# generate_ppt.py – No JSON + Blob Template Based
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64

from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image
from azure.storage.blob import BlobServiceClient

from utils import (
    get_env, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

# ------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------
ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)

AZURE_BLOB_CONN = get_env("AZURE_BLOB_CONN", required=True)
SOURCE_CONTAINER = get_env("AZURE_BLOB_CONTAINER", "ppt-dataset")


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    """Try to detect 'N slides' from user prompt."""
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        return int(match.group(1))
    return None


def fallback_plan(n):
    """Fallback slide plan if LLM fails."""
    slides = []
    for i in range(n):
        slides.append(
            {
                "title": f"Slide {i + 1}",
                "bullets": [
                    f"Key takeaway for slide {i + 1}",
                    f"Supporting detail {i + 1}.1",
                    f"Supporting detail {i + 1}.2",
                ],
            }
        )
    return slides


# ------------------------------------------------------------
# TEXT-BASED LLM (NO JSON)
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    """
    Ask LLM to return a text plan like:

    Slide 1: Title
    - Bullet
    - Bullet

    Slide 2: Title
    - Bullet
    ...
    """
    references_text = references_text or []

    sys_prompt = (
        "You are a professional presentation creator.\n\n"
        "Return the slide plan in this EXACT TEXT format only:\n\n"
        "Slide 1: Title\n"
        "- Bullet\n"
        "- Bullet\n"
        "- Bullet\n\n"
        "Slide 2: Title\n"
        "- Bullet\n"
        "- Bullet\n\n"
        f"You MUST create exactly {num_slides} slides.\n"
        "- Do NOT return JSON.\n"
        "- Do NOT add explanations.\n"
        "- Do NOT put multiple slides inside a single 'Slide X'.\n\n"
        "Use this content as reference:\n"
        f"{' '.join(references_text)[:2500]}"
    )

    user_prompt = f"Create a professional presentation for: {prompt}"

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt},
            ],
            max_completion_tokens=1400,
            temperature=0.7,
        )

        raw_text = resp.choices[0].message.content.strip()
        return parse_text_plan(raw_text, num_slides)

    except Exception as e:
        logger.warning(f"LLM plan failed → fallback used: {e}")
        return fallback_plan(num_slides)


# ------------------------------------------------------------
# TEXT → STRUCTURED SLIDES
# ------------------------------------------------------------
def parse_text_plan(text, num_slides):
    """
    Parse the LLM text output into a list of:
    {title: str, bullets: [str]}
    """
    slides = []

    # Split on "Slide X:" markers
    blocks = re.split(r"\n(?=Slide \d+:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        # Title line
        title_line = lines[0]
        # e.g., "Slide 1: Introduction"
        if ":" in title_line:
            title = title_line.split(":", 1)[1].strip()
        else:
            title = title_line.replace("Slide", "").strip()

        bullets = []
        for l in lines[1:]:
            if l.startswith("-"):
                bullets.append(l[1:].strip())

        # Enforce minimum bullets and cap max to avoid huge walls of text
        if len(bullets) < 3:
            bullets += [f"Additional point {i+1}" for i in range(3 - len(bullets))]
        bullets = bullets[:6]

        slides.append(
            {
                "title": title or "Untitled",
                "bullets": bullets,
            }
        )

    # If we didn't even get num_slides slides, signal "insufficient content"
    if len(slides) < num_slides:
        logger.warning("Parsed slides < requested slides; treating as insufficient content.")
        return None

    # Trim to exactly num_slides
    return slides[:num_slides]


# ------------------------------------------------------------
# IMAGE GENERATION (OPTIONAL)
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    """Generate an image from the GPT image model. Returns temp file path or None."""
    if not prompt:
        return None

    img_prompt = prompt + " Minimal professional illustration. No text."

    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=img_prompt,
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# TEMPLATE LOADER FROM AZURE BLOB
# ------------------------------------------------------------
def get_template_ppt_path(template_style: str | None) -> str | None:
    """
    If template_style is a blob name (not 'Plain (Default)'),
    download that PPT from SOURCE_CONTAINER to a temp file and return its path.
    Otherwise return None for plain Presentation().
    """
    if not template_style or template_style == "Plain (Default)":
        return None

    try:
        blob_service = BlobServiceClient.from_connection_string(AZURE_BLOB_CONN)
        container_client = blob_service.get_container_client(SOURCE_CONTAINER)

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
        with open(tmp.name, "wb") as f:
            stream = container_client.download_blob(template_style)
            stream.readinto(f)

        logger.info(f"Using blob '{template_style}' as PPT template.")
        return tmp.name
    except Exception as e:
        logger.warning(f"Failed to download template '{template_style}', falling back to plain: {e}")
        return None


# ------------------------------------------------------------
# PPT BUILDER (TEMPLATE-AWARE)
# ------------------------------------------------------------
def build_ppt(slides, template_style=None):
    template_path = get_template_ppt_path(template_style)

    if template_path:
        try:
            prs = Presentation(template_path)
        except Exception:
            logger.warning("Failed to open template PPT, falling back to blank.")
            prs = Presentation()
    else:
        prs = Presentation()

    for sp in slides:
        # Use Title + Content layout (index 1)
        slide = prs.slides.add_slide(prs.slide_layouts[1])

        # Title
        slide.shapes.title.text = sp.get("title", "")

        # Body
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        for b in sp.get("bullets", []):
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(20)

        body.top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.3)

        img_path = sp.get("image_path")

        if not img_path:
            # No image → full width for text
            body.left = Inches(0.5)
            body.width = prs.slide_width - Inches(1.0)
            continue

        # Image present → split layout
        body.left = Inches(0.5)
        body.width = prs.slide_width - Inches(4.0)

        try:
            img = Image.open(img_path)
            w, h = img.size
            aspect = w / h if h else 1.0

            max_w = Inches(3.0)
            max_h = Inches(2.5)

            if aspect >= 1:
                final_w = max_w
                final_h = final_w / aspect
            else:
                final_h = max_h
                final_w = final_h * aspect

            left = prs.slide_width - final_w - Inches(0.5)
            top = body.top

            slide.shapes.add_picture(img_path, left, top, width=final_w, height=final_h)
        except Exception:
            logger.exception("Image placement failed")

    out_path = os.path.join(
        tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx"
    )
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# MAIN PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt: str,
    requested_num_slides=None,
    template_style=None,
    image_required=False,
    tag_filters=None,
):
    """
    Main entry used by app.py.

    Returns:
        (ppt_path or None, log_dict)
        - If log_dict["error"] is True, ppt_path will be None and no PPT should be offered.
    """

    # 1) Semantic search in Chroma
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []

    if not refs:
        msg = (
            "I couldn’t find any relevant content in your uploaded presentations "
            "for this prompt. Please try a prompt related to the topics in your sample PPTs."
        )
        logger.warning(msg)
        return None, {"error": True, "message": msg}

    # 2) Build reference text snippets
    reference_text = []
    for r in refs:
        snippet = (r.get("text") or "")[:500]
        if snippet:
            reference_text.append(snippet)

    logger.info(f"Using {len(reference_text)} reference snippets from Chroma.")

    # 3) Slide count from prompt / UI
    detected_slides = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected_slides or 5

    # 4) Ask LLM for slide plan (text-based, no JSON)
    plan = call_llm_plan(
        prompt=prompt,
        references_text=reference_text,
        num_slides=num_slides,
    )

    if plan is None:
        msg = (
            "Not enough relevant content to generate this many slides. "
            "Try reducing the slide count or rephrasing your prompt."
        )
        logger.warning(msg)
        return None, {"error": True, "message": msg}

    # 5) Build slide objects (and optionally generate images)
    slides = []
    for sp in plan:
        img_path = None
        if image_required:
            img_path = generate_visual_image(sp.get("title"))

        slides.append(
            {
                "title": sp.get("title", "Untitled"),
                "bullets": sp.get("bullets", []),
                "image_path": img_path,
            }
        )

    # 6) Build PPT
    out_path = build_ppt(slides, template_style=template_style)

    # 7) Upload to Blob + log
    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides),
        "ppt_file": fname,
        "error": False,
        "template_style": template_style,
        "image_required": image_required,
    }

    upload_json_to_blob(
        json.dumps(log, indent=2).encode("utf-8"),
        f"logs/{fname}.json",
    )

    return out_path, log
