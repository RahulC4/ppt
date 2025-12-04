# ============================================================
# generate_ppt.py – GENUPDATE4 FINAL
# Clean PPT + Title + Agenda + Images + Design JSON Support
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.dml.color import RGBColor

from utils import (
    get_env, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

# ------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------
ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)

DESIGN_JSON_DIR = "design_jsons"

# ------------------------------------------------------------
# UTIL: LOAD DESIGN JSON
# ------------------------------------------------------------
def load_any_design_json():
    if not os.path.exists(DESIGN_JSON_DIR):
        return None

    files = os.listdir(DESIGN_JSON_DIR)
    if not files:
        return None

    try:
        with open(os.path.join(DESIGN_JSON_DIR, files[0]), "r", encoding="utf-8") as f:
            return json.load(f)
    except:
        return None


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    return int(match.group(1)) if match else None


def fallback_plan(n):
    return [
        {
            "title": f"Slide {i+1}",
            "bullets": [
                f"Key takeaway {i+1}",
                f"Supporting detail {i+1}.1",
                f"Supporting detail {i+1}.2",
            ],
        }
        for i in range(n)
    ]


# ------------------------------------------------------------
# LLM PLAN (NO JSON)
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    references_text = references_text or []

    sys_prompt = (
        "Return ONLY this format:\n"
        "Slide 1: Title\n"
        "- Bullet\n"
        "- Bullet\n"
        "- Bullet\n\n"
        f"Create EXACTLY {num_slides} slides."
    )

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": prompt},
            ],
            max_completion_tokens=1400,
            temperature=0.9,
        )

        raw = resp.choices[0].message.content.strip()
        return parse_text_plan(raw, num_slides)

    except Exception as e:
        logger.warning(f"LLM failed → fallback used: {e}")
        return fallback_plan(num_slides)


def parse_text_plan(text, num_slides):
    slides = []
    blocks = re.split(r"\n(?=Slide\s+\d+\s*:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        title = lines[0].split(":", 1)[-1].strip()
        bullets = [l[1:].strip() for l in lines[1:] if l.startswith("-")]
        bullets = bullets[:8]

        slides.append({"title": title, "bullets": bullets})

    return slides[:num_slides]


# ------------------------------------------------------------
# IMAGE GENERATION
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " professional minimal illustration, no text",
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# STYLE APPLY
# ------------------------------------------------------------
def apply_design(slide, design):
    if not design:
        return

    bg = design.get("theme_colors", [])
    fonts = design.get("font_families", [])

    try:
        if bg:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor.from_string(bg[0].replace("#", ""))
    except:
        pass

    try:
        if fonts:
            for shp in slide.shapes:
                if shp.has_text_frame:
                    for p in shp.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.name = fonts[0]
    except:
        pass


# ------------------------------------------------------------
# BULLET TEXTBOX
# ------------------------------------------------------------
def add_bullet_textbox(slide, bullets, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.word_wrap = True
    tf.clear()

    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(20)
        p.font.color.rgb = RGBColor(0, 0, 0)

    return tb


# ------------------------------------------------------------
# IMAGE PLACEMENT (ALIGNMENT PRESERVED)
# ------------------------------------------------------------
def add_image(slide, img_path, body):
    try:
        body.width = int(body.width * 0.6)
        left = body.left + body.width + Inches(0.3)
        top = body.top

        slide.shapes.add_picture(img_path, left, top, width=Inches(3))
    except:
        pass


# ------------------------------------------------------------
# FINAL PPT BUILDER
# ------------------------------------------------------------
def build_ppt(plan, image_required):
    design = load_any_design_json()
    prs = Presentation()

    # -------- TITLE SLIDE (NO IMAGE) --------
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = plan[0]["title"]
    subtitle = title_slide.placeholders[1]
    subtitle.text = datetime.now().strftime("%B %Y")
    apply_design(title_slide, design)

    # -------- AGENDA SLIDE (IMAGE ALLOWED) --------
    agenda_slide = prs.slides.add_slide(prs.slide_layouts[1])
    agenda_slide.shapes.title.text = "Agenda"

    agenda_items = [s["title"] for s in plan]

    body = add_bullet_textbox(
        agenda_slide,
        agenda_items,
        left=Inches(0.7),
        top=Inches(2),
        width=prs.slide_width - Inches(1.5),
        height=Inches(4),
    )

    if image_required:
        img = generate_visual_image("Agenda overview")
        if img:
            add_image(agenda_slide, img, body)

    apply_design(agenda_slide, design)

    # -------- CONTENT SLIDES --------
    for sp in plan:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = sp["title"]

        body = add_bullet_textbox(
            slide,
            sp["bullets"],
            left=Inches(0.7),
            top=Inches(2),
            width=prs.slide_width - Inches(1.5),
            height=Inches(4.5),
        )

        if image_required:
            img = generate_visual_image(sp["title"])
            if img:
                add_image(slide, img, body)

        apply_design(slide, design)

    out = os.path.join(tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out)
    return out


# ------------------------------------------------------------
# MAIN PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt,
    requested_num_slides=5,
    tag_filters=None,
    template_style=None,
    image_required=False,
):
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []

    reference_text = [(r.get("text") or "")[:500] for r in refs]

    detected = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected or 5

    plan = call_llm_plan(prompt, reference_text, num_slides)

    ppt_path = build_ppt(plan, image_required)
    total_slides = len(plan) + 2  # Title + Agenda

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(ppt_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": total_slides,
        "ppt_file": fname,
        "image_required": image_required,
        "template_style": template_style,
        "error": False,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode(), f"logs/{fname}.json")
    return ppt_path, log
