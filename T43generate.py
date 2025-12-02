# ============================================================
# generate_ppt.py – NO JSON | Images Checkbox | Safe Scaling
# WITH CORPORATE STYLING (NO ALIGNMENT CHANGES)
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image

from utils import (
    get_env, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        return int(match.group(1))
    return None


def fallback_plan(n):
    slides = []
    for i in range(n):
        slides.append({
            "title": f"Slide {i+1}",
            "bullets": [
                f"Key takeaway for slide {i+1}",
                f"Supporting detail {i+1}.1",
                f"Supporting detail {i+1}.2",
            ]
        })
    return slides


# ------------------------------------------------------------
# ✅ TEXT-BASED LLM (NO JSON)
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    references_text = references_text or []

    sys_prompt = (
        "You are a professional presentation creator.\n\n"
        "Return the slide plan in this EXACT TEXT format only:\n\n"
        "Slide 1: Title\n"
        "- Bullet\n"
        "- Bullet\n"
        "- Bullet\n\n"
        f"You MUST create exactly {num_slides} slides.\n"
        "Do NOT return JSON.\n"
        "Do NOT add explanations.\n\n"
        "Reference content:\n"
        f"{' '.join(references_text)[:2500]}"
    )

    user_prompt = f"Create a professional presentation for: {prompt}"

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt},
            ],
            max_completion_tokens=1400,
            temperature=1,
        )

        raw_text = resp.choices[0].message.content.strip()
        return parse_text_plan(raw_text, num_slides)

    except Exception as e:
        logger.warning(f"LLM failed → fallback used: {e}")
        return fallback_plan(num_slides)


# ------------------------------------------------------------
# ✅ TEXT → STRUCTURED SLIDES
# ------------------------------------------------------------
def parse_text_plan(text, num_slides):
    slides = []

    blocks = re.split(r"\n(?=Slide \d+:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        title_line = lines[0]
        title = title_line.replace("Slide", "").split(":", 1)[-1].strip()

        bullets = []
        for l in lines[1:]:
            if l.startswith("-"):
                bullets.append(l.replace("-", "").strip())

        if len(bullets) < 3:
            bullets += [f"Additional point {i+1}" for i in range(3 - len(bullets))]

        slides.append({
            "title": title,
            "bullets": bullets[:6],
        })

    if len(slides) < num_slides:
        return None

    return slides[:num_slides]


# ------------------------------------------------------------
# IMAGE GENERATION (OPTIONAL)
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " Professional minimal illustration. No text.",
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# ✅ CORPORATE STYLE (NO LAYOUT CHANGES)
# ------------------------------------------------------------
def apply_corporate_style(slide):
    try:
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(245, 247, 250)
    except:
        pass

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue

        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.color.rgb = RGBColor(32, 55, 100)

                if shape == slide.shapes.title:
                    run.font.size = Pt(28)
                    run.font.bold = True
                else:
                    run.font.size = Pt(20)


# ------------------------------------------------------------
# PPT BUILDER ✅ (ALIGNMENT UNCHANGED)
# ------------------------------------------------------------
def build_ppt(slides, template_style=None):
    prs = Presentation()

    for sp in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = sp["title"]

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        for b in sp["bullets"]:
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(20)

        body.top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.3)

        if sp.get("image_path"):
            try:
                body.width = prs.slide_width - Inches(4)
                slide.shapes.add_picture(
                    sp["image_path"],
                    prs.slide_width - Inches(3.5),
                    body.top,
                    width=Inches(3),
                )
            except:
                pass
        else:
            body.width = prs.slide_width - Inches(1)

        # ✅ APPLY CORPORATE STYLE SAFELY
        if template_style and template_style.lower() == "corporate":
            apply_corporate_style(slide)

    out_path = os.path.join(
        tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx"
    )
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# ✅ FINAL SAFE PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt,
    requested_num_slides=5,
    tag_filters=None,
    template_style=None,
    image_required=False,
):
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []

    if not refs:
        return None, {"error": True, "message": "No matching content found in sample PPTs."}

    reference_text = [(r.get("text") or "")[:500] for r in refs]

    detected_slides = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected_slides or 5

    plan = call_llm_plan(
        prompt=prompt,
        references_text=reference_text,
        num_slides=num_slides,
    )

    if not plan:
        return None, {
            "error": True,
            "message": "Not enough relevant content to generate this many slides. Try fewer slides or rephrase."
        }

    slides = []
    for sp in plan:
        img_path = generate_visual_image(sp["title"]) if image_required else None

        slides.append({
            "title": sp["title"],
            "bullets": sp["bullets"],
            "image_path": img_path,
        })

    out_path = build_ppt(slides, template_style=template_style)

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides),
        "ppt_file": fname,
        "error": False,
        "image_required": image_required,
        "template_style": template_style,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode(), f"logs/{fname}.json")

    return out_path, log
