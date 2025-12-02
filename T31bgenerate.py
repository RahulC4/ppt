# ============================================================
# generate_ppt.py – NO JSON | Images Checkbox | Safe Scaling
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import PP_PLACEHOLDER
from PIL import Image

from utils import (
    get_env, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        return int(match.group(1))
    return None


def fallback_plan(n):
    slides = []
    for i in range(n):
        slides.append({
            "title": f"Slide {i+1}",
            "bullets": [
                f"Key takeaway for slide {i+1}",
                f"Supporting detail {i+1}.1",
                f"Supporting detail {i+1}.2",
            ]
        })
    return slides


# ------------------------------------------------------------
# ✅ TEXT-BASED LLM (NO JSON)
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    references_text = references_text or []

    sys_prompt = (
        "You are a professional presentation creator.\n\n"
        "Return the slide plan in this EXACT TEXT format only:\n\n"
        "Slide 1: Title\n"
        "- Bullet\n"
        "- Bullet\n"
        "- Bullet\n\n"
        f"You MUST create exactly {num_slides} slides.\n"
        "Do NOT return JSON.\n"
        "Do NOT add explanations.\n\n"
        "Reference content:\n"
        f"{' '.join(references_text)[:2500]}"
    )

    user_prompt = f"Create a professional presentation for: {prompt}"

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt},
            ],
            max_completion_tokens=1400,
            temperature=0.7,
        )

        raw_text = resp.choices[0].message.content.strip()
        return parse_text_plan(raw_text, num_slides)

    except Exception as e:
        logger.warning(f"LLM failed → fallback used: {e}")
        return fallback_plan(num_slides)


# ------------------------------------------------------------
# ✅ TEXT → STRUCTURED SLIDES
# ------------------------------------------------------------
def parse_text_plan(text, num_slides):
    slides = []

    blocks = re.split(r"\n(?=Slide \d+:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        title_line = lines[0]
        title = title_line.replace("Slide", "").split(":", 1)[-1].strip()

        bullets = []
        for l in lines[1:]:
            if l.startswith("-"):
                bullets.append(l.replace("-", "").strip())

        # ✅ Enforce minimum bullets per slide
        if len(bullets) < 3:
            bullets += [f"Additional point {i+1}" for i in range(3 - len(bullets))]

        slides.append({
            "title": title,
            "bullets": bullets[:6],  # ✅ cap to avoid overflow
        })

    # ✅ Enforce slide count strictly
    if len(slides) < num_slides:
        return None  # 🚨 Signals "insufficient content"

    return slides[:num_slides]


# ------------------------------------------------------------
# IMAGE GENERATION (OPTIONAL)
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " Professional minimal illustration. No text.",
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# PPT BUILDER – USE TITLE+CONTENT LAYOUT, RESPECT THEME BOX
# ------------------------------------------------------------
def build_ppt(slides):
    prs = Presentation()

    # Standard "Title and Content" layout – keeps theme's box, bullets, etc.
    layout_idx = 1 if len(prs.slide_layouts) > 1 else 0
    content_layout = prs.slide_layouts[layout_idx]

    for sp in slides:
        slide = prs.slides.add_slide(content_layout)

        # ---- Title in the layout's title placeholder ----
        try:
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                slide.shapes.title.text = sp["title"]
        except Exception:
            # fall back to any TITLE placeholder
            for ph in slide.placeholders:
                if ph.placeholder_format.type == PP_PLACEHOLDER.TITLE and getattr(ph, "has_text_frame", False):
                    ph.text = sp["title"]
                    break

        # ---- Find the BODY placeholder for bullets ----
        body = None
        for ph in slide.placeholders:
            if ph.placeholder_format.type == PP_PLACEHOLDER.BODY and getattr(ph, "has_text_frame", False):
                body = ph
                break

        if body is None:
            # fallback: first non-title text shape
            for shp in slide.shapes:
                if getattr(shp, "has_text_frame", False) and shp is not slide.shapes.title:
                    body = shp
                    break

        if body is not None:
            tf = body.text_frame
            tf.clear()

            first = True
            for b in sp["bullets"]:
                if first:
                    p = tf.paragraphs[0]
                    first = False
                else:
                    p = tf.add_paragraph()
                p.text = b
                p.level = 0
                p.font.size = Pt(18)
        else:
            logger.warning("No BODY placeholder found on slide; skipping bullets")

        # ---- Optional image, docked inside the content area ----
        if sp.get("image_path"):
            try:
                img = Image.open(sp["image_path"])
                w, h = img.size
                aspect = w / h if h else 1.0

                # target width inside the body box
                max_width = Inches(2.5)
                img_width = max_width
                img_height = img_width / aspect

                if body is not None:
                    # Place at bottom-right inside the same content box
                    left = body.left + body.width - img_width
                    top = body.top + body.height - img_height
                else:
                    # fallback to bottom-right of slide
                    left = prs.slide_width - img_width - Inches(0.5)
                    top = prs.slide_height - img_height - Inches(0.5)

                slide.shapes.add_picture(sp["image_path"], left, top, width=img_width)

            except Exception:
                logger.exception("Image placement failed")

    out_path = os.path.join(
        tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx"
    )
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# ✅ FINAL SAFE PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt,
    requested_num_slides=5,
    tag_filters=None,
    template_style=None,
    image_required=False,
):
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []

    if not refs:
        return None, {"error": True, "message": "No matching content found in sample PPTs."}

    reference_text = [(r.get("text") or "")[:500] for r in refs]

    detected_slides = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected_slides or 5

    plan = call_llm_plan(
        prompt=prompt,
        references_text=reference_text,
        num_slides=num_slides,
    )

    # 🚨 HARD STOP if not enough content
    if not plan:
        return None, {
            "error": True,
            "message": "Not enough relevant content to generate this many slides. Try fewer slides or rephrase."
        }

    slides = []
    for sp in plan:
        img_path = generate_visual_image(sp["title"]) if image_required else None

        slides.append({
            "title": sp["title"],
            "bullets": sp["bullets"],
            "image_path": img_path,
        })

    out_path = build_ppt(slides)

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides),
        "ppt_file": fname,
        "error": False,
        "image_required": image_required,
        "template_style": template_style,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode(), f"logs/{fname}.json")

    return out_path, log
