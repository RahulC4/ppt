# ============================================================
# generate_ppt.py – NO JSON | Images Checkbox | Corporate Template
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import PP_PLACEHOLDER
from PIL import Image

from utils import (
    get_env, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)

# Where we expect template files to live
TEMPLATE_DIR = "templates"
TEMPLATE_MAP = {
    "Corporate": "corporate.pptx",  # templates/corporate.pptx
}

# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        return int(match.group(1))
    return None


def fallback_plan(n):
    slides = []
    for i in range(n):
        slides.append({
            "title": f"Slide {i+1}",
            "bullets": [
                f"Key takeaway for slide {i+1}",
                f"Supporting detail {i+1}.1",
                f"Supporting detail {i+1}.2",
            ]
        })
    return slides


def _set_title(slide, text: str):
    """Safely set the title text on any slide."""
    # Try proper title placeholders first
    for ph in slide.placeholders:
        if ph.placeholder_format.type in (
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE,
        ):
            ph.text = text
            return

    # Fallback: first text frame
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            shape.text_frame.text = text
            return


def _set_bullets(slide, bullets):
    """Safely set body bullets on any slide."""
    body_shape = None

    # Prefer BODY / CONTENT placeholders
    for ph in slide.placeholders:
        if ph.placeholder_format.type in (
            PP_PLACEHOLDER.BODY,
            PP_PLACEHOLDER.CONTENT,
        ):
            body_shape = ph
            break

    # Fallback: any text shape that is not the title
    if body_shape is None:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                # Try not to reuse the title shape
                if shape is slide.shapes.title:
                    continue
                body_shape = shape
                break

    if body_shape is None:
        # Nothing we can reasonably write into
        logger.warning("No suitable body placeholder found on slide; skipping bullets.")
        return

    tf = body_shape.text_frame
    tf.clear()

    first = True
    for b in bullets:
        if first and tf.paragraphs:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(20)


def _update_year_on_slide(slide):
    """Replace 20xx year numbers on a slide with the current year."""
    current_year = str(datetime.now().year)
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        tf = shape.text_frame
        old = "\n".join(p.text for p in tf.paragraphs)
        new = re.sub(r"\b20\d{2}\b", current_year, old)
        if new != old:
            tf.text = new


# ------------------------------------------------------------
# ✅ TEXT-BASED LLM (NO JSON)
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    references_text = references_text or []

    sys_prompt = (
        "You are a professional presentation creator.\n\n"
        "Return the slide plan in this EXACT TEXT format only:\n\n"
        "Slide 1: Title\n"
        "- Bullet\n"
        "- Bullet\n"
        "- Bullet\n\n"
        f"You MUST create exactly {num_slides} slides.\n"
        "Do NOT return JSON.\n"
        "Do NOT add explanations.\n\n"
        "Reference content:\n"
        f\"{' '.join(references_text)[:2500]}\"
    )

    user_prompt = f"Create a professional presentation for: {prompt}"

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt},
            ],
            max_completion_tokens=1400,
            temperature=0.7,
        )

        raw_text = resp.choices[0].message.content.strip()
        return parse_text_plan(raw_text, num_slides)

    except Exception as e:
        logger.warning(f"LLM failed → fallback used: {e}")
        return fallback_plan(num_slides)


# ------------------------------------------------------------
# ✅ TEXT → STRUCTURED SLIDES
# ------------------------------------------------------------
def parse_text_plan(text, num_slides):
    slides = []

    blocks = re.split(r"\n(?=Slide \d+:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        title_line = lines[0]
        # e.g. "Slide 1: Some title"
        if ":" in title_line:
            title = title_line.split(":", 1)[1].strip()
        else:
            title = title_line

        bullets = []
        for l in lines[1:]:
            if l.startswith("-"):
                bullets.append(l.replace("-", "").strip())

        # ✅ Enforce minimum bullets per slide
        if len(bullets) < 3:
            bullets += [f"Additional point {i+1}" for i in range(3 - len(bullets))]

        slides.append({
            "title": title,
            "bullets": bullets[:6],  # cap bullets to avoid overflow
        })

    # ✅ Enforce slide count strictly
    if len(slides) < num_slides:
        return None  # signal insufficient content

    return slides[:num_slides]


# ------------------------------------------------------------
# IMAGE GENERATION (OPTIONAL)
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " Professional minimal illustration. No text.",
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# PPT BUILDERS
# ------------------------------------------------------------
def build_plain_ppt(slides):
    """Original plain layout builder (no corporate template)."""
    prs = Presentation()

    for sp in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title + content
        slide.shapes.title.text = sp["title"]

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        for b in sp["bullets"]:
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(20)

        body.top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.3)

        if sp.get("image_path"):
            try:
                body.width = prs.slide_width - Inches(4)
                slide.shapes.add_picture(
                    sp["image_path"],
                    prs.slide_width - Inches(3.5),
                    body.top,
                    width=Inches(3),
                )
            except Exception:
                logger.exception("Image placement failed")

    out_path = os.path.join(tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out_path)
    return out_path


def build_corporate_ppt(slides, template_style: str):
    """Use templates/corporate.pptx as a style source and rebuild deck.

    Final structure:
    1. Corporate title slide (title = slides[0]['title'], year updated)
    2. Corporate agenda slide (bullets = all slide titles)
    3..N+2: content slides for each item in `slides`
    Last: Corporate thank-you slide (year updated)
    """
    template_file = TEMPLATE_MAP.get(template_style)
    if not template_file:
        logger.warning(f"No template mapping for style {template_style}; falling back to plain PPT.")
        return build_plain_ppt(slides)

    template_path = os.path.join(TEMPLATE_DIR, template_file)
    if not os.path.exists(template_path):
        logger.warning(f"Template file not found: {template_path}; falling back to plain PPT.")
        return build_plain_ppt(slides)

    prs = Presentation(template_path)

    if len(prs.slides) < 2:
        logger.warning("Corporate template has fewer than 2 slides; falling back to plain PPT.")
        return build_plain_ppt(slides)

    # Capture layouts from the existing slides before we wipe them
    title_layout = prs.slides[0].slide_layout
    agenda_layout = prs.slides[1].slide_layout
    thank_layout = prs.slides[-1].slide_layout

    # Use agenda layout as generic content layout as well
    content_layout = agenda_layout

    # Remove all existing slides but keep masters / layouts
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        sldIdLst.remove(sldId)

    # ---- 1. Title slide ----
    title_slide = prs.slides.add_slide(title_layout)
    main_title = slides[0]["title"] if slides else "Presentation Overview"
    _set_title(title_slide, main_title)
    _update_year_on_slide(title_slide)

    # ---- 2. Agenda slide ----
    agenda_slide = prs.slides.add_slide(agenda_layout)
    agenda_bullets = [s["title"] for s in slides] if slides else ["Overview"]
    _set_title(agenda_slide, "Agenda")
    _set_bullets(agenda_slide, agenda_bullets)

    # ---- 3..N+2: content slides ----
    for sp in slides:
        slide = prs.slides.add_slide(content_layout)
        _set_title(slide, sp["title"])
        _set_bullets(slide, sp["bullets"])

        # Images: **do not** add on title/agenda/thank-you; only here.
        if sp.get("image_path"):
            try:
                # try to put image on the right side
                body_shape = None
                for ph in slide.placeholders:
                    if ph.placeholder_format.type in (
                        PP_PLACEHOLDER.BODY,
                        PP_PLACEHOLDER.CONTENT,
                    ):
                        body_shape = ph
                        break
                if body_shape is not None:
                    body = body_shape
                    body.top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.3)
                    body.width = prs.slide_width - Inches(4)

                slide.shapes.add_picture(
                    sp["image_path"],
                    prs.slide_width - Inches(3.5),
                    slide.shapes.title.top + slide.shapes.title.height + Inches(0.5),
                    width=Inches(3),
                )
            except Exception:
                logger.exception("Image placement failed on corporate slide")

    # ---- Last: Thank-you slide ----
    thank_slide = prs.slides.add_slide(thank_layout)
    _update_year_on_slide(thank_slide)

    out_path = os.path.join(tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# ✅ FINAL SAFE PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt,
    requested_num_slides=5,
    tag_filters=None,
    template_style=None,
    image_required=False,
):
    # 1. Retrieve content from Chroma
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []
    if not refs:
        return None, {"error": True, "message": "No matching content found in sample PPTs."}

    reference_text = [(r.get("text") or "")[:500] for r in refs]

    # 2. Decide slide count
    detected_slides = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected_slides or 5

    # 3. Build slide plan from LLM
    plan = call_llm_plan(
        prompt=prompt,
        references_text=reference_text,
        num_slides=num_slides,
    )

    if not plan:
        return None, {
            "error": True,
            "message": (
                "Not enough relevant content to generate this many slides. "
                "Try fewer slides or rephrase the prompt."
            ),
        }

    # 4. Attach images if required
    slides = []
    for sp in plan:
        img_path = generate_visual_image(sp["title"]) if image_required else None
        slides.append(
            {
                "title": sp["title"],
                "bullets": sp["bullets"],
                "image_path": img_path,
            }
        )

    # 5. Build PPT – corporate template or plain
    if template_style and template_style.lower() == "corporate":
        out_path = build_corporate_ppt(slides, template_style="Corporate")
    else:
        out_path = build_plain_ppt(slides)

    # 6. Upload + log
    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides) + (3 if template_style and template_style.lower() == "corporate" else 0),
        "ppt_file": fname,
        "error": False,
        "image_required": image_required,
        "template_style": template_style,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode(), f"logs/{fname}.json")

    return out_path, log
