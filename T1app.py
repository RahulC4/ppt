# ============================================================
# generate_ppt.py – FINAL STABLE VERSION
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image

from utils import (
    get_env, safe_json_load, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob


# ------------------------------------------------------------
# NORMALIZE PPT NAME
# ------------------------------------------------------------
def normalize_name(path):
    if not path:
        return ""
    return path.replace("\\", "/").split("/")[-1]


# ------------------------------------------------------------
# INITIAL SETUP
# ------------------------------------------------------------
ensure_dir("design_jsons")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)


# ------------------------------------------------------------
# GAMMA-LIKE FIXED TEMPLATE PRESETS
# ------------------------------------------------------------
TEMPLATE_PRESETS = {
    "Plant": {
        "background_color": "#F7FAF0",
        "title_font": "Calibri",
        "body_font": "Calibri",
    },
    "Dark": {
        "background_color": "#1F2933",
        "title_font": "Segoe UI",
        "body_font": "Segoe UI",
    },
    "Minimal": {
        "background_color": "#FFFFFF",
        "title_font": "Arial",
        "body_font": "Arial",
    },
    "Corporate": {
        "background_color": "#F3F4F6",
        "title_font": "Calibri",
        "body_font": "Calibri",
    }
}


# ------------------------------------------------------------
# USER INTENT PARSER
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    num_slides = None
    theme = None

    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        num_slides = int(match.group(1))

    for t in ["corporate", "modern", "minimal", "professional", "dark", "light"]:
        if t in prompt.lower():
            theme = t.capitalize()
            break

    return num_slides, theme


# ------------------------------------------------------------
# GPT PLAN GENERATOR (REFERENCE-ONLY)
# ------------------------------------------------------------
def call_llm_plan(
    prompt,
    style,
    design_context,
    references_text,
    num_slides=None,
    theme=None,
    text_density="Concise",
):

    density_rules = {
        "Minimal": "Use at most 1–2 short bullet points per slide.",
        "Concise": "Use about 3 bullet points per slide.",
        "Detailed": "Use about 5 bullet points per slide.",
        "Extensive": "Use 6–8 detailed bullet points per slide.",
    }

    density_text = density_rules.get(text_density, "")

    sys_prompt = (
        "You are a senior presentation designer.\n"
        "You MUST use ONLY the provided reference snippets.\n"
        "Do NOT invent new topics.\n\n"
        "Return STRICT JSON ONLY in this format:\n"
        "[{\"title\": str, \"bullets\": [str], \"visual_required\": bool, "
        "\"visual_prompt\": str }]\n\n"
        "If the user asks for images → set visual_required=true.\n"
        "Do NOT include text inside image prompts.\n"
        f"{density_text}\n\n"
        "REFERENCE SNIPPETS:\n"
        f"{json.dumps(references_text)[:2500]}"
    )

    user_prompt = f"Create a {style} presentation plan for: {prompt}"
    if num_slides:
        user_prompt += f". Use exactly {num_slides} slides."

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt},
            ],
            max_completion_tokens=1500,
            temperature=0.9,
        )

        plan = safe_json_load(resp.choices[0].message.content)

        if not plan:
            logger.warning("Invalid plan JSON. Using fallback.")
            return [{"title": "Intro", "bullets": ["Overview"], "visual_required": False}]

        return plan

    except Exception:
        logger.exception("Plan generation failed")
        return [{"title": "Intro", "bullets": ["Overview"], "visual_required": False}]


# ------------------------------------------------------------
# IMAGE GENERATION
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    img_prompt = (prompt or "") + " Minimal professional illustration. No text labels."

    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=img_prompt,
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if b64:
            img_bytes = base64.b64decode(b64)
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
            tmp.write(img_bytes)
            tmp.close()
            return tmp.name

        url = getattr(resp.data[0], "url", None)
        if url:
            import requests
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
            tmp.write(requests.get(url).content)
            tmp.close()
            return tmp.name

        return None

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# COLOR PARSER
# ------------------------------------------------------------
def parse_hex_color(hex_str):
    try:
        if not hex_str:
            return None
        hex_str = hex_str.replace("#", "")
        r = int(hex_str[0:2], 16)
        g = int(hex_str[2:4], 16)
        b = int(hex_str[4:6], 16)
        return RGBColor(r, g, b)
    except:
        return None


# ------------------------------------------------------------
# DESIGN JSON MATCHING
# ------------------------------------------------------------
def extract_design_for_slide(design_context, matched_ppt_name, matched_slide_idx):

    norm_target = normalize_name(matched_ppt_name)

    for design in design_context:
        norm_json = normalize_name(design["ppt_name"])

        if norm_json == norm_target:
            for slide in design["slides"]:
                if slide["index"] == matched_slide_idx:
                    return slide

    return None


# ------------------------------------------------------------
# APPLY DESIGN STYLE (AUTO + PRESET)
# ------------------------------------------------------------
def apply_design_style(slide, prs, design_meta, template_style="Auto"):

    preset = TEMPLATE_PRESETS.get(template_style)

    # 1️⃣ Apply fixed UI template
    if preset:
        bg = parse_hex_color(preset.get("background_color"))
        if bg:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = bg

        try:
            title_font = preset.get("title_font")
            body_font = preset.get("body_font")

            if title_font:
                for r in slide.shapes.title.text_frame.paragraphs[0].runs:
                    r.font.name = title_font

            if body_font:
                body = slide.placeholders[1]
                for p in body.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.name = body_font
        except:
            pass

    # 2️⃣ Apply extracted design ONLY if template_style == Auto
    if not design_meta or template_style != "Auto":
        return

    text_fonts = design_meta.get("text_fonts", [])
    if text_fonts:
        try:
            for p in slide.placeholders[1].text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = text_fonts[0]
        except:
            pass

    bg = parse_hex_color(design_meta.get("background_color"))
    if bg:
        try:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = bg
        except:
            pass


# ------------------------------------------------------------
# PPT BUILDER (FULL WIDTH FIX)
# ------------------------------------------------------------
def build_ppt(slides, matched_designs, template_style="Auto"):
    prs = Presentation()

    for idx, s in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])

        apply_design_style(slide, prs, matched_designs[idx], template_style)

        slide.shapes.title.text = s.get("title", "")

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        for b in s.get("bullets", []):
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(18)

        body.top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.3)

        # ✅ NO IMAGE → FULL WIDTH TEXT
        if not s.get("image_path"):
            body.left = Inches(0.5)
            body.width = prs.slide_width - Inches(1.0)
            body.height = prs.slide_height - Inches(1.2)
            tf.word_wrap = True
            continue

        # ✅ IMAGE EXISTS → TEXT LEFT, IMAGE RIGHT
        body.left = Inches(0.5)
        body.width = prs.slide_width - Inches(4.0)

        try:
            img_path = s["image_path"]
            img = Image.open(img_path)
            w, h = img.size
            aspect = w / h

            max_w = Inches(3.0)
            max_h = Inches(2.8)

            if aspect >= 1:
                final_w = max_w
                final_h = final_w / aspect
            else:
                final_h = max_h
                final_w = final_h * aspect

            left = prs.slide_width - final_w - Inches(0.4)
            top = body.top

            slide.shapes.add_picture(img_path, left, top, width=final_w, height=final_h)

        except Exception:
            logger.exception("Image placement failed")

    out_path = os.path.join(tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# MAIN PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt: str,
    style="Auto",
    requested_num_slides=None,
    theme=None,
    tag_filters=None,
    text_density="Concise",
    template_style="Auto",
):

    detected_slides, detected_theme = parse_user_intent(prompt)
    requested_num_slides = requested_num_slides or detected_slides
    theme = theme or detected_theme

    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []

    # ✅ FRIENDLY NO-MATCH RESPONSE
    if not refs:
        msg = (
            "I couldn’t find any relevant content in your uploaded sample presentations "
            "for this request. Please rephrase your prompt using topics related to your "
            "existing PPTs, or upload a new sample PPT."
        )
        logger.warning(msg)
        return None, {"error": True, "message": msg}

    design_context = []
    reference_text = []

    for r in refs:
        ppt_name = r.get("ppt_name")
        snippet = (r.get("text") or "")[:500]
        if snippet:
            reference_text.append(snippet)

        json_path = os.path.join("design_jsons", normalize_name(ppt_name) + ".json")
        if os.path.exists(json_path):
            try:
                with open(json_path, "r", encoding="utf-8") as f:
                    design_context.append(json.load(f))
            except:
                pass

    logger.info(f"Loaded {len(design_context)} design JSONs and {len(reference_text)} text snippets.")

    plan = call_llm_plan(
        prompt,
        style,
        design_context,
        reference_text,
        num_slides=requested_num_slides,
        theme=theme,
        text_density=text_density,
    )

    force_images = "image" in prompt.lower() or "images" in prompt.lower()

    slides = []
    matched_designs = []

    for sp in plan:

        best = refs[0]
        try:
            matched_idx = int(best.get("slide_index"))
        except:
            matched_idx = 0

        design_meta = extract_design_for_slide(
            design_context,
            best.get("ppt_name"),
            matched_idx,
        )

        matched_designs.append(design_meta)

        img = None
        if sp.get("visual_required") or force_images:
            img = generate_visual_image(sp.get("visual_prompt"))

        slides.append({
            "title": sp.get("title"),
            "bullets": sp.get("bullets", []),
            "image_path": img,
        })

    out_path = build_ppt(slides, matched_designs, template_style)

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides),
        "ppt_file": fname,
        "text_density": text_density,
        "template_style": template_style,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode(), f"logs/{fname}.json")

    return out_path, log
