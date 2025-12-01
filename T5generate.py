# ============================================================
# generate_ppt.py – FIXED + TEXT DENSITY + FIXED TEMPLATE MODE
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image
from utils import (
    get_env, safe_json_load, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob


# ------------------------------------------------------------
# INITIAL SETUP
# ------------------------------------------------------------

ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)


# ------------------------------------------------------------
# USER INTENT PARSER
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    num_slides = None
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        num_slides = int(match.group(1))
    return num_slides


# ------------------------------------------------------------
# GPT PLAN GENERATOR (FIXED ✅)
# ------------------------------------------------------------
def call_llm_plan(prompt, style,
                  design_context=None, references_text=None,
                  num_slides=None, theme=None,
                  text_density=None):

    sys_prompt = (
        "You are a presentation planner.\n"
        "Return STRICT JSON ONLY in this format:\n"
        "[{\"title\": str, \"bullets\": [str], "
        "\"visual_required\": bool, \"visual_prompt\": str }]\n"
        "If images are requested → set visual_required=true.\n"
        "Do NOT put text inside image prompts.\n"
        "Output valid JSON only."
    )

    user_prompt = f"Create a professional presentation plan for: {prompt}"

    if num_slides:
        user_prompt += f". Make exactly {num_slides} slides."

    if text_density:
        user_prompt += f". Use {text_density} amount of text per slide."

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt}
            ],
            max_completion_tokens=1200,
            temperature=1
        )

        plan = safe_json_load(resp.choices[0].message.content)

        if not plan or not isinstance(plan, list):
            raise ValueError("Invalid JSON plan")

        return plan

    except Exception as e:
        logger.warning("Invalid plan JSON, using fallback plan")

        # ✅ SAFE FALLBACK (NEVER EMPTY)
        slides = []
        n = num_slides or 5
        for i in range(n):
            slides.append({
                "title": f"Slide {i+1}",
                "bullets": ["Key point 1", "Key point 2"],
                "visual_required": False,
                "visual_prompt": ""
            })
        return slides


# ------------------------------------------------------------
# IMAGE GENERATION
# ------------------------------------------------------------
def generate_visual_image(prompt: str):

    if not prompt:
        return None

    img_prompt = prompt + " Minimal clean illustration. No text."

    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=img_prompt,
            size="1024x1024"
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if b64:
            img_bytes = base64.b64decode(b64)
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
            tmp.write(img_bytes)
            tmp.close()
            return tmp.name

        return None

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# FIXED TEMPLATE PPT BUILDER ✅
# ------------------------------------------------------------
def build_ppt(slides):
    prs = Presentation()

    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])

        # Title
        slide.shapes.title.text = s["title"]

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        for b in s.get("bullets", []):
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(20)

        # ✅ FULL WIDTH IF NO IMAGE
        if not s.get("image_path"):
            body.width = prs.slide_width - Inches(1)
            continue

        # ✅ IMAGE MODE (RIGHT SIDE)
        body.width = prs.slide_width - Inches(4)

        try:
            img = Image.open(s["image_path"])
            w, h = img.size
            aspect = w / h

            max_w = Inches(3)
            max_h = Inches(2.5)

            if aspect >= 1:
                final_w = max_w
                final_h = final_w / aspect
            else:
                final_h = max_h
                final_w = final_h * aspect

            left = prs.slide_width - final_w - Inches(0.5)
            top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.2)

            slide.shapes.add_picture(s["image_path"], left, top,
                                     width=final_w, height=final_h)

        except:
            logger.exception("Image placement failed")

    out_path = os.path.join(
        tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx"
    )
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# ✅ ✅ ✅ MAIN PIPELINE (CLEAN + SAFE)
# ------------------------------------------------------------
def generate_presentation(prompt: str,
                          style="Auto",
                          requested_num_slides=None,
                          theme=None,
                          text_density="concise",
                          tag_filters=None):

    # ✅ STEP 1: Semantic Search
    refs = semantic_search(prompt, top_k=5, tags=tag_filters)

    if not refs:
        return None, {
            "error": "No matching content found. Please use a prompt related to uploaded presentations."
        }

    # ✅ STEP 2: Prompt → Plan
    detected_slides = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected_slides or 5

    plan = call_llm_plan(
        prompt=prompt,
        style=style,
        num_slides=num_slides,
        text_density=text_density
    )

    # ✅ STEP 3: Image Handling
    force_images = "image" in prompt.lower() or "images" in prompt.lower()

    slides = []
    for sp in plan:
        img = None
        if sp.get("visual_required") or force_images:
            img = generate_visual_image(sp.get("visual_prompt"))

        slides.append({
            "title": sp.get("title"),
            "bullets": sp.get("bullets", []),
            "image_path": img
        })

    # ✅ STEP 4: Build PPT
    out_path = build_ppt(slides)

    # ✅ STEP 5: Upload
    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides),
        "ppt_file": fname
    }
    upload_json_to_blob(json.dumps(log, indent=2).encode(), f"logs/{fname}.json")

    return out_path, log
