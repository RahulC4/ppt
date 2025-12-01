# ============================================================
# generate_ppt.py – Semantic + Fixed Layout + UI Friendly
# (robust JSON parsing from content only)
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64

from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image

from utils import (
    get_env, safe_json_load, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

# ------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------
ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    """Try to detect 'N slides' from user prompt."""
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        return int(match.group(1))
    return None


def _fallback_plan(n: int | None):
    """Pure Python fallback so we never crash."""
    n = n or 3
    fb = []
    for i in range(n):
        fb.append(
            {
                "title": f"Slide {i + 1}",
                "bullets": [
                    f"Key point 1 for slide {i + 1}",
                    f"Key point 2 for slide {i + 1}",
                ],
                "visual_required": False,
                "visual_prompt": "",
            }
        )
    return fb


def _extract_json_like(raw: str):
    """
    Try to rescue JSON from a response that may contain extra text.
    Returns a Python object or None.
    """
    if not raw:
        return None

    # First try direct parse
    parsed = safe_json_load(raw)
    if parsed is not None:
        return parsed

    # Try to find a JSON array
    start = raw.find("[")
    end = raw.rfind("]")
    if start != -1 and end != -1 and end > start:
        candidate = raw[start : end + 1]
        parsed = safe_json_load(candidate)
        if parsed is not None:
            return parsed

    # Try to find a JSON object
    start = raw.find("{")
    end = raw.rfind("}")
    if start != -1 and end != -1 and end > start:
        candidate = raw[start : end + 1]
        parsed = safe_json_load(candidate)
        if parsed is not None:
            return parsed

    return None


# ------------------------------------------------------------
# LLM PLAN GENERATOR (JSON-enforced, robust)
# ------------------------------------------------------------
def call_llm_plan(
    prompt,
    style,
    design_context=None,
    references_text=None,
    num_slides=None,
    theme=None,
    text_density=None,  # kept for signature compatibility, NOT used inside
):
    """
    Ask GPT to create a slide plan.
    Returns a list of dicts: {title, bullets, visual_required, visual_prompt}.
    If anything goes wrong, we fall back to a simple deterministic plan.
    ALWAYS returns a list whose length == num_slides (if num_slides is given).
    """
    references_text = references_text or []

    sys_prompt = (
        "You are a presentation planner.\n"
        "You MUST respond with JSON ONLY.\n\n"
        "Preferred shape:\n"
        "{\n"
        '  \"slides\": [\n'
        '    {\n'
        '      \"title\": \"string\",\n'
        '      \"bullets\": [\"string\", \"string\", ...],\n'
        '      \"visual_required\": true or false,\n'
        '      \"visual_prompt\": \"string\"\n'
        "    }, ...\n"
        "  ]\n"
        "}\n\n"
        "It is ALSO acceptable to respond as a bare JSON array:\n"
        "[{...}, {...}, ...]\n\n"
        "Rules:\n"
        "- No markdown, no comments, no text outside the JSON.\n"
        "- Every bullet must be a plain string (no bullet symbols or numbering).\n"
        "- If the user asks for images, set visual_required=true.\n"
        "- Do NOT put any written text or labels inside the generated images.\n\n"
        "You may use these reference snippets as guidance:\n"
        f"{json.dumps(references_text)[:2000]}"
    )

    user_prompt = f"Create a professional presentation plan for: {prompt}"
    if num_slides:
        user_prompt += f" Use exactly {num_slides} slides."

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt},
            ],
            max_completion_tokens=1200,
            temperature=1,
        )

        raw = (resp.choices[0].message.content or "").strip()
        parsed = _extract_json_like(raw)

        if isinstance(parsed, dict):
            slides_raw = parsed.get("slides")
        elif isinstance(parsed, list):
            slides_raw = parsed
        else:
            raise ValueError("Could not parse valid JSON slide plan")

        if not isinstance(slides_raw, list) or not slides_raw:
            raise ValueError("slides list missing or empty")

        # Clean & normalize items
        cleaned = []
        for item in slides_raw:
            if not isinstance(item, dict):
                continue
            cleaned.append(
                {
                    "title": item.get("title", "Untitled"),
                    "bullets": item.get("bullets", []),
                    "visual_required": bool(item.get("visual_required", False)),
                    "visual_prompt": item.get("visual_prompt", ""),
                }
            )

        if not cleaned:
            raise ValueError("Cleaned slides list is empty")

        plan = cleaned

        # --- ENFORCE SLIDE COUNT ---
        if num_slides and num_slides > 0:
            if len(plan) > num_slides:
                plan = plan[:num_slides]
            elif len(plan) < num_slides:
                # duplicate last slide as "(cont.)" until we reach num_slides
                while len(plan) < num_slides:
                    last = dict(plan[-1])
                    last["title"] = last.get("title", "Slide") + " (cont.)"
                    plan.append(last)

        return plan

    except Exception as e:
        logger.warning(f"Invalid plan JSON, using fallback plan: {e}")
        return _fallback_plan(num_slides)


# ------------------------------------------------------------
# IMAGE GENERATION
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    """Generate an image file from GPT image model. Returns a temp file path or None."""
    if not prompt:
        return None

    img_prompt = prompt + " Minimal, professional illustration. No text labels."

    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=img_prompt,
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# PPT BUILDER (FULL-WIDTH TEXT WHEN NO IMAGE)
# ------------------------------------------------------------
def build_ppt(slides):
    prs = Presentation()

    for sp in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])

        # Title
        slide.shapes.title.text = sp.get("title", "")

        # Body placeholder
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        for b in sp.get("bullets", []):
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(20)

        # Space under title
        body.top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.3)

        img_path = sp.get("image_path")

        if not img_path:
            # NO IMAGE → full width
            body.left = Inches(0.5)
            body.width = prs.slide_width - Inches(1.0)
            continue

        # IMAGE PRESENT → split layout
        body.left = Inches(0.5)
        body.width = prs.slide_width - Inches(4.0)

        try:
            img = Image.open(img_path)
            w, h = img.size
            aspect = w / h if h else 1.0

            max_w = Inches(3.0)
            max_h = Inches(2.5)

            if aspect >= 1:
                final_w = max_w
                final_h = final_w / aspect
            else:
                final_h = max_h
                final_w = final_h * aspect

            left = prs.slide_width - final_w - Inches(0.5)
            top = body.top

            slide.shapes.add_picture(img_path, left, top, width=final_w, height=final_h)
        except Exception:
            logger.exception("Image placement failed")

    out_path = os.path.join(
        tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx"
    )
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# MAIN PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt: str,
    style="Auto",
    requested_num_slides=None,
    theme=None,
    text_density="Concise",
    tag_filters=None,
    template_style=None,  # accepted for UI compatibility, not used yet
):
    """
    Main entry used by app.py / test scripts.

    Returns:
        (ppt_path or None, log_dict)
        - If log_dict["error"] is True, ppt_path will be None and no PPT should be offered.
    """

    # 1) Semantic search in Chroma
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []

    if not refs:
        msg = (
            "I couldn’t find any relevant content in your uploaded presentations "
            "for this prompt. Please try a prompt related to the topics in your sample PPTs."
        )
        logger.warning(msg)
        return None, {"error": True, "message": msg}

    # 2) Build reference text snippets
    reference_text = []
    for r in refs:
        snippet = (r.get("text") or "")[:500]
        if snippet:
            reference_text.append(snippet)

    logger.info(f"Using {len(reference_text)} reference snippets from Chroma.")

    # 3) Slide count from prompt / UI
    detected_slides = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected_slides or 5

    # 4) Ask LLM for slide plan (density is applied LATER)
    plan = call_llm_plan(
        prompt=prompt,
        style=style,
        references_text=reference_text,
        num_slides=num_slides,
        theme=theme,
        text_density=None,
    )

    # ---------- POST-PROCESS BULLETS BASED ON text_density ----------
    density_to_max_bullets = {
        "Minimal": 2,
        "Concise": 3,
        "Detailed": 5,
        "Extensive": 7,
    }
    max_bullets = density_to_max_bullets.get(text_density, 3)
    MAX_CHARS = 160

    for s in plan:
        # clean title
        s["title"] = (s.get("title") or "Untitled").strip()[:80]

        raw_bullets = s.get("bullets") or []
        cleaned_bullets = []
        for b in raw_bullets[:max_bullets]:
            t = str(b).strip()
            if not t:
                continue
            if len(t) > MAX_CHARS:
                t = t[:MAX_CHARS] + "..."
            cleaned_bullets.append(t)
        s["bullets"] = cleaned_bullets

    # Validate plan: if everything is empty, bail out with a friendly message
    if not plan or not isinstance(plan, list):
        msg = (
            "I couldn't generate a valid slide plan from your prompt and the "
            "available sample PPTs. Please simplify or rephrase your request."
        )
        logger.warning(msg)
        return None, {"error": True, "message": msg}

    all_empty = True
    for s in plan:
        title = (s.get("title") or "").strip()
        bullets = [str(b).strip() for b in (s.get("bullets") or []) if str(b).strip()]
        if title or bullets:
            all_empty = False
            break

    if all_empty:
        msg = (
            "I generated only empty slides for this request. "
            "Please try a clearer prompt or one that is closer to your sample PPT content."
        )
        logger.warning(msg)
        return None, {"error": True, "message": msg}

    # 5) Build slide objects (and optionally generate images)
    force_images = "image" in prompt.lower() or "images" in prompt.lower()
    slides = []

    for sp in plan:
        img_path = None
        if sp.get("visual_required") or force_images:
            img_path = generate_visual_image(sp.get("visual_prompt"))

        slides.append(
            {
                "title": sp.get("title", "Untitled"),
                "bullets": sp.get("bullets", []),
                "image_path": img_path,
            }
        )

    # 6) Build PPT
    out_path = build_ppt(slides)

    # 7) Upload to Blob + log
    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides),
        "ppt_file": fname,
        "error": False,
        "text_density": text_density,
        "template_style": template_style,
    }

    upload_json_to_blob(
        json.dumps(log, indent=2).encode("utf-8"),
        f"logs/{fname}.json",
    )

    return out_path, log
