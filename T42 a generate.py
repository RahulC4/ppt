# ============================================================
# generate_ppt.py – NO JSON | Images Checkbox | Corporate Template | AUTO SIZE FIX
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE

from utils import (
    get_env,
    logger,
    now_ts,
    ensure_dir,
    text_client,
    image_client,
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

# ------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------
ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
CORPORATE_TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "corporate.pptx")


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    return int(match.group(1)) if match else None


def fallback_plan(n):
    slides = []
    for i in range(n):
        slides.append({
            "title": f"Slide {i+1}",
            "bullets": [
                f"Main point {i+1}",
                f"Supporting detail {i+1}.1",
                f"Supporting detail {i+1}.2",
            ],
        })
    return slides


# ------------------------------------------------------------
# ✅ LLM PLAN (NO JSON)
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    references_text = references_text or []

    sys_prompt = (
        "Return slide plan in this EXACT format:\n\n"
        "Slide 1: Title\n- Bullet\n- Bullet\n- Bullet\n\n"
        f"You MUST create exactly {num_slides} slides.\n"
        "DO NOT RETURN JSON.\n\n"
        f"Reference content:\n{' '.join(references_text)[:2500]}"
    )

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": prompt},
            ],
            max_completion_tokens=1400,
            temperature=0.7,
        )
        raw = resp.choices[0].message.content.strip()
        return parse_text_plan(raw, num_slides)
    except Exception as e:
        logger.warning(f"LLM failed → fallback used: {e}")
        return fallback_plan(num_slides)


def parse_text_plan(text, num_slides):
    slides = []
    blocks = re.split(r"\n(?=Slide\s+\d+\s*:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        title = lines[0].split(":", 1)[-1].strip()
        bullets = [l[1:].strip() for l in lines[1:] if l.startswith("-")]

        if len(bullets) < 3:
            bullets += [f"Additional point {i+1}" for i in range(3 - len(bullets))]

        slides.append({"title": title, "bullets": bullets[:8]})

    return slides[:num_slides] if len(slides) >= num_slides else None


# ------------------------------------------------------------
# IMAGE GENERATION
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " professional minimal illustration, no text",
            size="1024x1024",
        )
        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name
    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# ✅✅✅ CORE CONTENT WRITER (GENUPDATE4 LOGIC)
# ------------------------------------------------------------
def render_slide_content(slide, title, bullets, image_path=None):
    slide.shapes.title.text = title
    body = slide.placeholders[1]
    tf = body.text_frame

    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.word_wrap = True
    tf.clear()

    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(20)

    body.top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.3)

    if image_path:
        body.width = slide.part.slide_layout.part.slide_width - Inches(4)
        body.left = Inches(0.5)
        slide.shapes.add_picture(
            image_path,
            slide.part.slide_layout.part.slide_width - Inches(3.5),
            body.top,
            width=Inches(3),
        )
    else:
        body.left = Inches(0.5)
        body.width = slide.part.slide_layout.part.slide_width - Inches(1)


# ------------------------------------------------------------
# DEFAULT BUILDER
# ------------------------------------------------------------
def build_default_ppt(plan, image_required):
    prs = Presentation()

    for sp in plan:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        img = generate_visual_image(sp["title"]) if image_required else None
        render_slide_content(slide, sp["title"], sp["bullets"], img)

    out = os.path.join(tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out)
    return out


# ------------------------------------------------------------
# CORPORATE BUILDER (GENUPDATE3 STRUCTURE)
# ------------------------------------------------------------
def build_corporate_ppt(plan, image_required):
    prs = Presentation(CORPORATE_TEMPLATE_PATH)

    title_layout = prs.slides[0].slide_layout
    agenda_layout = prs.slides[1].slide_layout
    content_layout = prs.slides[2].slide_layout
    thanks_layout = prs.slides[-1].slide_layout

    while prs.slides:
        xml_slides = prs.slides._sldIdLst
        prs.part.drop_rel(xml_slides[0].rId)
        del xml_slides[0]

    # Title
    title_slide = prs.slides.add_slide(title_layout)
    title_slide.shapes.title.text = plan[0]["title"]
    for shp in title_slide.shapes:
        if shp.has_text_frame and shp is not title_slide.shapes.title:
            shp.text = datetime.now().strftime("%B %Y")

    # Agenda
    agenda_slide = prs.slides.add_slide(agenda_layout)
    agenda_slide.shapes.title.text = "Agenda"
    render_slide_content(
        agenda_slide,
        "Agenda",
        [s["title"] for s in plan],
        generate_visual_image("Agenda") if image_required else None,
    )

    # Content
    for sp in plan:
        slide = prs.slides.add_slide(content_layout)
        img = generate_visual_image(sp["title"]) if image_required else None
        render_slide_content(slide, sp["title"], sp["bullets"], img)

    # Thank You
    prs.slides.add_slide(thanks_layout)

    out = os.path.join(tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out)
    return out


# ------------------------------------------------------------
# ✅ FINAL PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt,
    requested_num_slides=5,
    tag_filters=None,
    template_style=None,
    image_required=False,
):
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []
    if not refs:
        return None, {"error": True, "message": "No matching content found in sample PPTs."}

    reference_text = [(r.get("text") or "")[:500] for r in refs]
    num_slides = requested_num_slides or parse_user_intent(prompt) or 5

    plan = call_llm_plan(prompt, reference_text, num_slides)
    if not plan:
        return None, {"error": True, "message": "Not enough content to generate slides."}

    if (template_style or "").lower() == "corporate":
        ppt_path = build_corporate_ppt(plan, image_required)
        total_slides = len(plan) + 3
    else:
        ppt_path = build_default_ppt(plan, image_required)
        total_slides = len(plan)

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(ppt_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": total_slides,
        "ppt_file": fname,
        "image_required": image_required,
        "template_style": template_style,
        "error": False,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode(), f"logs/{fname}.json")
    return ppt_path, log
