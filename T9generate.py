# ============================================================
# generate_ppt.py – AUTO | Title + Agenda + Thank You | Safe Scaling
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
from PIL import Image

from utils import (
    get_env, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        return int(match.group(1))
    return None


def fallback_plan(n):
    slides = []
    for i in range(n):
        slides.append({
            "title": f"Slide {i+1}",
            "bullets": [
                f"Key takeaway for slide {i+1}",
                f"Supporting detail {i+1}.1",
                f"Supporting detail {i+1}.2",
            ]
        })
    return slides


# ------------------------------------------------------------
# ✅ TEXT-BASED LLM
# ------------------------------------------------------------
def call_llm_plan_auto(prompt, references_text=None, num_slides=5):
    references_text = references_text or []

    sys_prompt = (
        "You are a professional presentation creator.\n\n"
        "Return the slide plan in this EXACT TEXT format only:\n\n"
        "Slide 1: Title\n"
        "- Bullet\n"
        "- Bullet\n"
        "- Bullet\n\n"
        f"You MUST create exactly {num_slides} slides.\n"
        "Do NOT return JSON.\n"
        "Do NOT add explanations.\n\n"
        "Reference content:\n"
        f"{' '.join(references_text)[:2500]}"
    )

    user_prompt = f"Create a professional presentation for: {prompt}"

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt},
            ],
            max_completion_tokens=1400,
            temperature=1,
        )

        raw_text = resp.choices[0].message.content.strip()
        return parse_text_plan(raw_text, num_slides)

    except Exception as e:
        logger.warning(f"LLM failed → fallback used: {e}")
        return fallback_plan(num_slides)


# ------------------------------------------------------------
# ✅ TEXT → STRUCTURED
# ------------------------------------------------------------
def parse_text_plan(text, num_slides):
    slides = []
    blocks = re.split(r"\n(?=Slide \d+:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        title_line = lines[0]
        title = title_line.replace("Slide", "").split(":", 1)[-1].strip()

        bullets = [l.replace("-", "").strip() for l in lines[1:] if l.startswith("-")]

        if len(bullets) < 3:
            bullets += [f"Additional point {i+1}" for i in range(3 - len(bullets))]

        slides.append({"title": title, "bullets": bullets[:6]})

    return slides[:num_slides]


# ------------------------------------------------------------
# ✅ IMAGE GENERATION
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " Professional minimal illustration. No text.",
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except Exception:
        return None


# ------------------------------------------------------------
# ✅ PPT BUILDER (TITLE + AGENDA + CONTENT + THANK YOU)
# ------------------------------------------------------------
def build_ppt(slides, agenda_titles, image_required):
    prs = Presentation()

    # ✅ TITLE SLIDE
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = slides[0]["title"]
    subtitle = title_slide.placeholders[1]
    subtitle.text = datetime.now().strftime("%B %Y")

    # ✅ AGENDA SLIDE
    agenda_slide = prs.slides.add_slide(prs.slide_layouts[1])
    agenda_slide.shapes.title.text = "Agenda"
    body = agenda_slide.placeholders[1]
    tf = body.text_frame
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.word_wrap = True
    tf.clear()

    for t in agenda_titles:
        p = tf.add_paragraph()
        p.text = t
        p.font.size = Pt(20)
        p.level = 0

    body.top = agenda_slide.shapes.title.top + agenda_slide.shapes.title.height + Inches(0.3)

    # ✅ Agenda Image (Allowed)
    if image_required:
        img = generate_visual_image("Agenda illustration")
        if img:
            body.width = prs.slide_width - Inches(4)
            body.left = Inches(0.5)
            agenda_slide.shapes.add_picture(
                img,
                prs.slide_width - Inches(3.5),
                body.top,
                width=Inches(3),
            )
    else:
            body.left = Inches(0.5)
            body.width = prs.slide_width - Inches(1)

    # ✅ CONTENT SLIDES (UNCHANGED LOGIC)
    for sp in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = sp["title"]

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.word_wrap = True
        tf.clear()

        for b in sp["bullets"]:
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(20)
            p.level = 0

        body.top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.3)

        if sp.get("image_path"):
            body.width = prs.slide_width - Inches(4)
            body.left = Inches(0.5)

            slide.shapes.add_picture(
                sp["image_path"],
                prs.slide_width - Inches(3.5),
                body.top,
                width=Inches(3),
            )
        else:
            body.left = Inches(0.5)
            body.width = prs.slide_width - Inches(1)

    # ✅ THANK YOU SLIDE (NO IMAGE)
    thank_slide = prs.slides.add_slide(prs.slide_layouts[5])
    thank_slide.shapes.title.text = "Thank You"

    out_path = os.path.join(tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# ✅ FINAL PIPELINE
# ------------------------------------------------------------
def generate_presentation_auto(
    prompt,
    requested_num_slides=5,
    tag_filters=None,
    template_style=None,
    image_required=False,
):
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []

    reference_text = [(r.get("text") or "")[:500] for r in refs]

    detected_slides = parse_user_intent(prompt)
    num_slides = requested_num_slides or detected_slides or 5

    plan = call_llm_plan_auto(prompt, reference_text, num_slides)

    slides = []
    for sp in plan:
        img_path = generate_visual_image(sp["title"]) if image_required else None
        slides.append({
            "title": sp["title"],
            "bullets": sp["bullets"],
            "image_path": img_path,
        })

    agenda_titles = [s["title"] for s in slides]

    out_path = build_ppt(slides, agenda_titles, image_required)

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides),
        "ppt_file": fname,
        "error": False,
        "image_required": image_required,
        "template_style": template_style,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode(), f"logs/{fname}.json")

    return out_path, log
