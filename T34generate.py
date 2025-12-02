# ============================================================
# generate_ppt.py – Corporate Template Engine (Final Stable)
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from utils import get_env, logger, now_ts, ensure_dir, text_client, image_client
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

ensure_dir("generated")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)

TEMPLATE_PATH = os.path.join("templates", "corporate.pptx")

# ------------------------------------------------------------
# PARSE USER SLIDE COUNT FROM PROMPT (OPTIONAL)
# ------------------------------------------------------------
def parse_user_intent(prompt: str):
    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        return int(match.group(1))
    return None


# ------------------------------------------------------------
# FALLBACK PLAN
# ------------------------------------------------------------
def fallback_plan(n):
    slides = []
    for i in range(n):
        slides.append({
            "title": f"Slide {i+1}",
            "bullets": [
                f"Key point {i+1}.1",
                f"Key point {i+1}.2",
                f"Key point {i+1}.3",
            ]
        })
    return slides


# ------------------------------------------------------------
# TEXT BASED LLM (NO JSON)
# ------------------------------------------------------------
def call_llm_plan(prompt, references_text=None, num_slides=5):
    references_text = references_text or []

    sys_prompt = (
        "Return slides in EXACT format:\n\n"
        "Slide 1: Title\n"
        "- Bullet\n"
        "- Bullet\n"
        "- Bullet\n\n"
        f"Create exactly {num_slides} slides.\n"
        "Do NOT return JSON.\n"
    )

    user_prompt = f"Create a professional presentation:\n{prompt}"

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt},
            ],
            max_completion_tokens=1400,
            temperature=0.7,
        )

        raw_text = resp.choices[0].message.content.strip()
        return parse_text_plan(raw_text, num_slides)

    except Exception as e:
        logger.warning(f"LLM failed → fallback used: {e}")
        return fallback_plan(num_slides)


# ------------------------------------------------------------
# TEXT → STRUCTURED SLIDES
# ------------------------------------------------------------
def parse_text_plan(text, num_slides):
    slides = []
    blocks = re.split(r"\n(?=Slide \d+:)", text)

    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue

        title = lines[0].split(":", 1)[-1].strip()
        bullets = [l.replace("-", "").strip() for l in lines[1:] if l.startswith("-")]

        if len(bullets) < 3:
            bullets += [f"Extra point {i+1}" for i in range(3 - len(bullets))]

        slides.append({"title": title, "bullets": bullets[:6]})

    if len(slides) < num_slides:
        return None

    return slides[:num_slides]


# ------------------------------------------------------------
# IMAGE GENERATION (ONLY FOR NEW SLIDES)
# ------------------------------------------------------------
def generate_visual_image(prompt: str):
    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " professional illustration",
            size="1024x1024",
        )

        b64 = getattr(resp.data[0], "b64_json", None)
        if not b64:
            return None

        img_bytes = base64.b64decode(b64)
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()
        return tmp.name

    except:
        return None


# ------------------------------------------------------------
# ✅ MASTER TEMPLATE PPT BUILDER
# ------------------------------------------------------------
def build_from_corporate_template(title_text, agenda_bullets, generated_slides, image_required):
    prs = Presentation(TEMPLATE_PATH)

    # ✅ Keep only 3 template slides (Title, Agenda, Thank You)
    while len(prs.slides) > 3:
        rId = prs.slides._sldIdLst[-1].rId
        prs.part.drop_rel(rId)
        prs.slides._sldIdLst.remove(prs.slides._sldIdLst[-1])

    # ✅ Update current year automatically
    current_year = str(datetime.now().year)

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                shape.text = shape.text.replace("2021", current_year)
                shape.text = shape.text.replace("2022", current_year)
                shape.text = shape.text.replace("2023", current_year)

    # ✅ Slide 1 – Title
    title_slide = prs.slides[0]
    for shape in title_slide.shapes:
        if shape.has_text_frame:
            shape.text = title_text

    # ✅ Slide 2 – Agenda
    agenda_slide = prs.slides[1]
    for shape in agenda_slide.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            for bullet in agenda_bullets:
                p = tf.add_paragraph()
                p.text = bullet

    # ✅ Generate new content slides
    content_layout = prs.slide_layouts[1]

    for sp in generated_slides:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = sp["title"]

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        for b in sp["bullets"]:
            p = tf.add_paragraph()
            p.text = b
            p.font.size = Pt(18)

        if image_required:
            img_path = generate_visual_image(sp["title"])
            if img_path:
                slide.shapes.add_picture(img_path, prs.slide_width - Inches(3.5), Inches(1), width=Inches(3))

    out_path = os.path.join(tempfile.gettempdir(), f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# ✅ FINAL PIPELINE
# ------------------------------------------------------------
def generate_presentation(
    prompt,
    requested_num_slides=5,
    tag_filters=None,
    template_style=None,
    image_required=False,
):
    refs = semantic_search(prompt, top_k=5, tags=tag_filters) or []

    if not refs:
        return None, {"error": True, "message": "No matching content found in sample PPTs."}

    detected_slides = parse_user_intent(prompt)
    num_generated_slides = requested_num_slides or detected_slides or 5

    reference_text = [(r.get("text") or "")[:500] for r in refs]

    plan = call_llm_plan(prompt, reference_text, num_generated_slides)

    if not plan:
        return None, {"error": True, "message": "Not enough relevant content."}

    title_text = plan[0]["title"]
    agenda_bullets = plan[0]["bullets"]

    out_path = build_from_corporate_template(
        title_text=title_text,
        agenda_bullets=agenda_bullets,
        generated_slides=plan,
        image_required=image_required,
    )

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "user_requested_slides": num_generated_slides,
        "final_slide_count": num_generated_slides + 3,
        "template_used": "corporate.pptx",
        "ppt_file": fname,
        "image_required": image_required,
        "error": False,
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode(), f"logs/{fname}.json")

    return out_path, log
