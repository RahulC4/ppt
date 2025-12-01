# ============================================================
# generate_ppt.py – FINAL STABLE TEMPLATE-BASED VERSION
# ============================================================

import os
import tempfile
import uuid
import json
import re
import base64
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image

from utils import (
    get_env, safe_json_load, logger, now_ts,
    ensure_dir, text_client, image_client
)
from search_utils import semantic_search
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob


# ------------------------------------------------------------
# ENV + SETUP
# ------------------------------------------------------------

ensure_dir("design_jsons")

CHAT_MODEL = get_env("CHAT_MODEL", required=True)
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)


# ------------------------------------------------------------
# USER INTENT PARSER
# ------------------------------------------------------------

def parse_user_intent(prompt: str):
    num_slides = None

    match = re.search(r"(\d+)\s+slides?", prompt.lower())
    if match:
        num_slides = int(match.group(1))

    return num_slides


# ------------------------------------------------------------
# TEMPLATE STYLE ENGINE (✅ REAL THEMES)
# ------------------------------------------------------------

def apply_template_style(slide, template_style):

    bg = slide.background.fill
    bg.solid()

    if template_style == "dark":
        bg.fore_color.rgb = RGBColor(20, 20, 20)
        title_color = RGBColor(255, 255, 255)
        body_color = RGBColor(220, 220, 220)

    elif template_style == "corporate":
        bg.fore_color.rgb = RGBColor(245, 246, 250)
        title_color = RGBColor(32, 55, 100)
        body_color = RGBColor(60, 60, 60)

    elif template_style == "modern":
        bg.fore_color.rgb = RGBColor(15, 76, 129)
        title_color = RGBColor(255, 255, 255)
        body_color = RGBColor(230, 230, 230)

    else:
        return

    # Title styling
    if slide.shapes.title:
        for p in slide.shapes.title.text_frame.paragraphs:
            for r in p.runs:
                r.font.color.rgb = title_color
                r.font.size = Pt(32)
                r.font.bold = True

    # Body styling
    body = slide.placeholders[1]
    for p in body.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = body_color
            r.font.size = Pt(20)


# ------------------------------------------------------------
# SAFE LLM PLAN GENERATOR
# ------------------------------------------------------------

def call_llm_plan(prompt, style, references_text, num_slides=None):

    sys_prompt = (
        "You are a presentation assistant.\n"
        "Return STRICT JSON ONLY:\n"
        "[{\"title\": str, \"bullets\": [str], \"visual_required\": bool, \"visual_prompt\": str}]\n"
        "Do NOT return markdown. Do NOT return explanations."
    )

    user_prompt = f"Create a {style} presentation plan ONLY from this content:\n\n"
    user_prompt += "\n".join(references_text)

    if num_slides:
        user_prompt += f"\n\nMake exactly {num_slides} slides."

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": user_prompt}
            ],
            max_completion_tokens=1200,
            temperature=1
        )

        plan = safe_json_load(resp.choices[0].message.content)

        if not plan:
            raise ValueError("Invalid JSON")

        return plan

    except Exception:
        logger.exception("LLM plan failed — using fallback")

        # ✅ NEVER EMPTY FALLBACK
        return [{
            "title": "Overview",
            "bullets": references_text[:3] if references_text else ["Content not found."],
            "visual_required": False,
            "visual_prompt": ""
        }]


# ------------------------------------------------------------
# IMAGE GENERATION
# ------------------------------------------------------------

def generate_visual_image(prompt: str):

    if not prompt:
        return None

    try:
        resp = image_client.images.generate(
            model=IMAGE_MODEL,
            prompt=prompt + " professional business illustration, no text",
            size="1024x1024"
        )

        img_b64 = resp.data[0].b64_json
        img_bytes = base64.b64decode(img_b64)

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
        tmp.write(img_bytes)
        tmp.close()

        return tmp.name

    except Exception:
        logger.exception("Image generation failed")
        return None


# ------------------------------------------------------------
# PPT BUILDER (✅ FULL WIDTH FIX + TEMPLATE)
# ------------------------------------------------------------

def build_ppt(slides, template_style):

    prs = Presentation()

    for s in slides:

        slide = prs.slides.add_slide(prs.slide_layouts[1])

        # ✅ APPLY TEMPLATE
        apply_template_style(slide, template_style)

        # Title
        slide.shapes.title.text = s["title"]

        # Body
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()

        for b in s["bullets"]:
            p = tf.add_paragraph()
            p.text = b

        body.top = Inches(1.5)
        body.left = Inches(0.7)

        # ✅ FULL WIDTH IF NO IMAGE
        if not s.get("image_path"):
            body.width = prs.slide_width - Inches(1.4)
            continue

        # ✅ SHRINK TEXT IF IMAGE EXISTS
        body.width = prs.slide_width - Inches(4.5)

        try:
            img = Image.open(s["image_path"])
            w, h = img.size
            aspect = w / h

            max_w = Inches(3.5)
            max_h = Inches(3)

            if aspect > 1:
                final_w = max_w
                final_h = max_w / aspect
            else:
                final_h = max_h
                final_w = max_h * aspect

            left = prs.slide_width - final_w - Inches(0.5)
            top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.2)

            slide.shapes.add_picture(
                s["image_path"], left, top,
                width=final_w, height=final_h
            )

        except Exception:
            logger.exception("Image placement failed")

    out_path = os.path.join(tempfile.gettempdir(),
                            f"generated_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(out_path)
    return out_path


# ------------------------------------------------------------
# ✅ FINAL MAIN PIPELINE (CHROMA GROUNDED ONLY)
# ------------------------------------------------------------

def generate_presentation(prompt: str,
                          template_style="corporate",
                          requested_num_slides=None):

    detected_slides = parse_user_intent(prompt)
    requested_num_slides = requested_num_slides or detected_slides or 5

    # ✅ SEMANTIC SEARCH (GROUNDING IS MANDATORY)
    refs = semantic_search(prompt, top_k=5)

    if not refs:
        return None, {
            "status": "no_match",
            "message": "No matching knowledge found. Please refine your prompt."
        }

    reference_text = [r["text"] for r in refs if r.get("text")]

    # ✅ LLM ONLY RESTRUCTURES FOUND CONTENT
    plan = call_llm_plan(
        prompt,
        template_style,
        reference_text,
        num_slides=requested_num_slides
    )

    slides = []

    for sp in plan:
        img = None
        if sp.get("visual_required"):
            img = generate_visual_image(sp.get("visual_prompt"))

        slides.append({
            "title": sp.get("title", "Slide"),
            "bullets": sp.get("bullets", []),
            "image_path": img
        })

    out_path = build_ppt(slides, template_style)

    fname = f"generated_{uuid.uuid4().hex[:8]}.pptx"
    upload_ppt_to_blob(out_path, fname)

    log = {
        "timestamp": now_ts(),
        "prompt": prompt,
        "slides_generated": len(slides),
        "template": template_style,
        "ppt_file": fname
    }

    upload_json_to_blob(json.dumps(log, indent=2).encode(),
                        f"logs/{fname}.json")

    return out_path, log
